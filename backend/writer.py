"""把汇总结果写入新建工作簿，或复制并填充指定的 Excel 模板。"""
from __future__ import annotations

import datetime as dt
import re
import shutil
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import db
from .utils import UserError, coerce_number, display_value, parse_date_value, safe_filename

BRAND = "0F6B5C"
BRAND_LIGHT = "E8F4F0"
LINE = "DDE6E1"
ZEBRA = "F7FAF9"
INK = "1D2A25"

THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10.5, color=INK)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
# 样式对象做成模块级常量：以前每写一个单元格就 new 一个 Alignment，
# 5 万行 × 6 列就是 30 万个对象，白占内存也拖慢写入。
HEADER_ALIGN = Alignment(horizontal="left", vertical="center")
BODY_ALIGN = Alignment(vertical="center", wrap_text=False)
WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)


def _safe_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", name or "汇总").strip() or "汇总"
    return name[:31]


def _as_cell_value(value: Any, col_type: str):
    if value is None:
        return None
    if col_type == "date":
        ok, parsed = parse_date_value(value)
        if ok and isinstance(parsed, (dt.date, dt.datetime)):
            return parsed
        # 解析不成日期：数字原样保留，别把 5 位工号 10001 变成文本 "10001"
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return value
        return display_value(value)
    if col_type == "number":
        num = coerce_number(value)
        return num if num is not None else display_value(value)
    if col_type == "text":
        return value if isinstance(value, str) else display_value(value)
    # auto：保留原生数字 / 日期
    if isinstance(value, (dt.date, dt.datetime, int, float, bool)):
        return value
    return display_value(value)


def _date_format() -> str:
    fmt = str(db.get_settings().get("date_format") or "yyyy-mm-dd")
    return {"yyyy-mm-dd": "yyyy-mm-dd", "yyyy/mm/dd": "yyyy/mm/dd",
            "yyyy年mm月dd日": "yyyy年mm月dd日"}.get(fmt, "yyyy-mm-dd")


def _is_percent_source(value: Any) -> bool:
    """源值是否是 "50%" 这类百分比文本（coerce_number 会把它折成 0.5）。"""
    return isinstance(value, str) and value.strip().endswith("%")


def _apply_number_format(cell, col_type: str, value: Any, percent: bool = False):
    if percent and isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.number_format = "0.00%"
    elif col_type == "date" and isinstance(value, (dt.date, dt.datetime)):
        cell.number_format = _date_format()
    elif col_type == "number" and isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.number_format = "#,##0.##"
    elif isinstance(value, dt.datetime):
        cell.number_format = "yyyy-mm-dd hh:mm"
    elif isinstance(value, dt.date):
        cell.number_format = _date_format()


def _display_len(value: Any) -> int:
    text = display_value(value)
    length = 0
    for ch in text:
        length += 2 if ord(ch) > 127 else 1
    return length


def _auto_widths(headers: list[str], rows: list[list[Any]], configured: dict[str, int]) -> list[int]:
    widths = []
    for i, header in enumerate(headers):
        if configured.get(header):
            widths.append(max(8, min(60, int(configured[header]))))
            continue
        longest = _display_len(header)
        for row in rows[:400]:
            if i < len(row):
                longest = max(longest, _display_len(row[i]))
        widths.append(max(10, min(42, longest + 3)))
    return widths


def _blank_workbook(output: Path, headers: list[str], rows: list[list[Any]],
                    config: dict, result: dict, profile_name: str) -> dict:
    wb = Workbook()
    ws = wb.active
    sheet_name = _safe_sheet_name((config.get("target") or {}).get("sheet") or "汇总")
    ws.title = sheet_name
    ws.sheet_view.showGridLines = False

    columns = [c for c in (config.get("target") or {}).get("columns", [])
               if c.get("included", True) and str(c.get("name") or "").strip()]
    types = [c.get("type") or "auto" for c in columns]
    widths_conf = {c.get("name"): c.get("width") for c in columns if c.get("width")}

    for i, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 30

    for r, row in enumerate(rows, start=2):
        zebra = ZEBRA_FILL if r % 2 == 0 else None
        for c, value in enumerate(row, start=1):
            col_type = types[c - 1] if c - 1 < len(types) else "auto"
            cell_value = _as_cell_value(value, col_type)
            cell = ws.cell(row=r, column=c)
            _set_safe(cell, cell_value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = BODY_ALIGN
            _apply_number_format(cell, col_type, cell_value, _is_percent_source(value))
            if zebra is not None:
                cell.fill = zebra

    for i, width in enumerate(_auto_widths(headers, rows, widths_conf), start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    if headers:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"

    # 汇总说明：让文件的来源、口径可追溯
    info = wb.create_sheet("汇总说明")
    info.sheet_view.showGridLines = False
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 72
    info["A1"] = "表格汇总助手 · 数据汇总说明"
    info["A1"].font = Font(name="微软雅黑", size=14, bold=True, color=BRAND)
    info.merge_cells("A1:B1")
    info_rows = [
        ("方案名称", profile_name),
        ("生成时间", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("汇总方式", "主表关联补齐" if config.get("mode") == "join" else "多表纵向合并"),
        ("数据规模", f"{result['stats']['rows']} 行 × {len(headers)} 列"),
        ("参与表格", "；".join(f"{s['label']}（{s['rows']} 行）" for s in result.get("sources", []))),
        ("空值单元格", str(result["stats"].get("empty_cells", 0))),
    ]
    for i, (k, v) in enumerate(info_rows, start=3):
        info.cell(row=i, column=1, value=k).font = Font(name="微软雅黑", size=10, bold=True, color="55635D")
        info.cell(row=i, column=2, value=v).font = Font(name="微软雅黑", size=10, color=INK)
        info.cell(row=i, column=2).alignment = WRAP_ALIGN
    if result.get("warnings"):
        row = len(info_rows) + 4
        info.cell(row=row, column=1, value="提示").font = Font(name="微软雅黑", size=10, bold=True, color="A9761F")
        info.cell(row=row, column=2, value="\n".join(result["warnings"])).font = Font(name="微软雅黑", size=10, color="6B5A33")
        info.cell(row=row, column=2).alignment = WRAP_ALIGN

    wb.save(output)
    return {"sheet": sheet_name, "header_row": 1, "data_start": 2, "headers": headers,
            "columns": list(range(1, len(headers) + 1))}


def _find_last_data_row(ws: Worksheet, columns: list[int], start: int) -> int:
    """只统计真正有值的单元格，避免模板整列设了格式导致 max_row 虚高。"""
    last = start - 1
    cells = getattr(ws, "_cells", None)
    if isinstance(cells, dict):
        for (row, col), cell in cells.items():
            if row >= start and col in columns and cell.value is not None and str(cell.value).strip() != "":
                last = max(last, row)
        return last
    upper = min(ws.max_row, start + 50000)
    for row in range(upper, start - 1, -1):
        for col in columns:
            value = ws.cell(row=row, column=col).value
            if value is not None and str(value).strip() != "":
                return row
    return last


def _set_safe(cell, value):
    """写入值，并阻止以 = + - @ 开头的文本被 Excel 当成公式。

    前导空白/换行会被 Excel 忽略，攻击者可用 "\\t=cmd" 绕过前缀判断，
    所以先 lstrip 再判断；命中时用 quotePrefix 强制文本（比改 data_type 稳）。
    """
    cell.value = value
    if isinstance(value, str) and value:
        # 快路径：首字符既不是公式前缀也不是空白，就不可能被 Excel 当公式，
        # 不必为每个单元格都做一次 lstrip。
        if value[0] not in "=+-@ \t\r\n":
            return
        probe = value.lstrip(" \t\r\n")
        if probe.startswith(("=", "+", "-", "@")):
            try:
                cell.quotePrefix = True
            except Exception:
                pass
            try:
                cell.data_type = "s"
            except Exception:
                pass


def _template_workbook(template_path: Path, output: Path, headers: list[str],
                       rows: list[list[Any]], config: dict, result: dict) -> dict:
    target = config.get("target") or {}
    sheet_name = target.get("sheet")
    header_row = int(target.get("header_row") or 1)
    columns = [c for c in target.get("columns", [])
               if c.get("included", True) and str(c.get("name") or "").strip()]
    types = [c.get("type") or "auto" for c in columns]

    suffix = template_path.suffix.lower()
    keep_vba = suffix == ".xlsm"
    try:
        wb = load_workbook(template_path, keep_vba=keep_vba)
    except Exception as exc:
        raise UserError(f"无法打开模板文件：{exc}") from exc
    if sheet_name not in wb.sheetnames:
        raise UserError("模板中找不到所选工作表，可能文件已被修改。")
    ws = wb[sheet_name]

    # 目标列在模板中的真实列号：优先使用导入时记录的 template_col，其次按列名查找
    template_headers = []
    for c in range(1, ws.max_column + 1):
        template_headers.append(str(ws.cell(row=header_row, column=c).value or "").strip())
    used_cols: list[int] = []
    for i, header in enumerate(headers):
        col = columns[i].get("template_col")
        try:
            col = int(col) + 1 if col is not None else None
        except (TypeError, ValueError):
            col = None
        if col is None or col < 1 or col > ws.max_column:
            try:
                col = template_headers.index(header) + 1
            except ValueError:
                raise UserError(f"模板表头中找不到列「{header}」，请重新导入模板列。")
        if col in used_cols:
            raise UserError(f"目标列「{header}」对应了重复的模板列，请检查列配置。")
        used_cols.append(col)

    data_start = header_row + 1
    last_data = _find_last_data_row(ws, used_cols, data_start)
    append_mode = (target.get("write_mode") or "replace") == "append"
    all_last = _find_last_data_row(ws, list(range(1, ws.max_column + 1)), data_start) if append_mode else last_data
    start_row = (all_last + 1) if append_mode else data_start

    if not append_mode:
        clear_to = max(last_data, data_start + len(rows) - 1)
        for row in range(data_start, clear_to + 1):
            for col in used_cols:
                cell = ws.cell(row=row, column=col)
                try:
                    cell.value = None
                except AttributeError:
                    raise UserError(f"模板数据区存在合并单元格（第 {row} 行），请先在 Excel 中取消合并后再试。")

    anchor_row = data_start
    anchor_height = ws.row_dimensions[anchor_row].height
    old_max = ws.max_row
    for r_offset, row in enumerate(rows):
        r = start_row + r_offset
        if r > old_max:
            ws.row_dimensions[r].height = anchor_height
        for c_offset, value in enumerate(row):
            col = used_cols[c_offset]
            col_type = types[c_offset] if c_offset < len(types) else "auto"
            cell_value = _as_cell_value(value, col_type)
            cell = ws.cell(row=r, column=col)
            try:
                _set_safe(cell, cell_value)
            except AttributeError:
                raise UserError(f"模板数据区存在合并单元格（第 {r} 行第 {col} 列），请先取消合并。")
            _apply_number_format(cell, col_type, cell_value, _is_percent_source(value))
            if r > old_max:
                # 复制模板首行数据区的样式，让新增行保持原有观感
                try:
                    anchor_cell = ws.cell(row=anchor_row, column=col)
                    cell._style = copy(anchor_cell._style)
                except Exception:
                    pass
    wb.save(output)
    return {"sheet": sheet_name, "header_row": header_row, "headers": headers,
            "columns": used_cols, "data_start": start_row}


def write_output(config: dict, result: dict, run_dir: Path, profile_name: str,
                 template: dict | None = None) -> dict:
    run_dir.mkdir(parents=True, exist_ok=True)
    target = config.get("target") or {}
    mode = target.get("kind") or "blank"
    now = dt.datetime.now()
    stamp = now.strftime("%Y%m%d_%H%M%S") + f"_{now.microsecond // 1000:03d}"
    base = safe_filename(f"汇总_{profile_name}_{stamp}")
    if mode == "template":
        if not template:
            raise UserError("目标模板不存在。")
        template_path = Path(template["path"])
        suffix = template_path.suffix.lower()
        out_suffix = ".xlsm" if suffix == ".xlsm" else ".xlsx"
        output = run_dir / f"{base}{out_suffix}"
        shutil.copyfile(template_path, output)
        info = _template_workbook(template_path, output, result["headers"],
                                  result["rows"], config, result)
    else:
        output = run_dir / f"{base}.xlsx"
        info = _blank_workbook(output, result["headers"], result["rows"],
                               config, result, profile_name)
    return {
        "path": str(output),
        "name": output.name,
        "size": output.stat().st_size,
        "sheet": info.get("sheet"),
        "header_row": info.get("header_row"),
        # data_start 用于数据对比定位"这次真正写入的第一行"。
        # append 模式下模板原有行在此之前，不能算进本次结果。
        "data_start": info.get("data_start") or int(info.get("header_row") or 1) + 1,
        "headers": result["headers"],
        "columns": info.get("columns") or list(range(1, len(result["headers"]) + 1)),
    }
