"""数据对比：按主键比较两次汇总结果 / 两张数据表，并导出差异 Excel。"""
from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import db
from .config import COMPARE_DIR
from .engine import Registry
from .utils import UserError, display_value, norm_key, now_iso, uid

MAX_CHANGE_ROWS = 800
MAX_ROW_PREVIEW = 300

BRAND = "0F6B5C"
LINE = "DDE6E1"
THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_ADD = PatternFill("solid", fgColor="E9F6EF")
FILL_DEL = PatternFill("solid", fgColor="FBECEA")
FILL_CHG = PatternFill("solid", fgColor="FDF4E3")
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10.5, color="1D2A25")


def load_side(spec: dict, registry: Registry | None = None) -> dict:
    kind = spec.get("type")
    if kind == "run":
        run = db.query_one("SELECT * FROM runs WHERE id=?", (spec.get("id"),))
        if not run:
            raise UserError("找不到要对比的运行记录。")
        if run.get("status") != "ok":
            raise UserError("这次运行没有成功生成结果，无法对比。")
        summary = db.loads(run.get("summary"), {}) or {}
        output = summary.get("output") or {}
        path = Path(run.get("output_path") or "")
        if not path.exists():
            raise UserError("结果文件不存在，可能已经被清理。")
        headers = output.get("headers") or []
        columns = output.get("columns") or []
        sheet = output.get("sheet") or "汇总"
        header_row = int(output.get("header_row") or 1)
        # append 模式下模板原有行在 data_start 之前，不能算作本次结果
        data_start = int(output.get("data_start") or header_row + 1)
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            if sheet not in wb.sheetnames:
                raise UserError("结果文件中找不到汇总工作表。")
            ws = wb[sheet]
            actual_headers = []
            for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row)):
                actual_headers.append(display_value(cell.value).strip())
            if not headers:
                headers = actual_headers
            if not columns:
                columns = [actual_headers.index(h) + 1 for h in headers if h in actual_headers]
            if len(columns) != len(headers):
                raise UserError("结果文件的列结构与记录不一致。")
            rows = []
            for row in ws.iter_rows(min_row=data_start, values_only=True):
                if not any(v is not None and str(v).strip() != "" for v in row):
                    continue
                rows.append([row[c - 1] if c - 1 < len(row) else None for c in columns])
            wb.close()
        except UserError:
            raise
        except Exception as exc:
            raise UserError(f"读取结果文件失败：{exc}") from exc
        return {
            "name": f"{run.get('profile_name') or '汇总结果'} · {run.get('created_at', '')[:16].replace('T', ' ')}",
            "headers": [str(h) for h in headers],
            "rows": rows,
            "kind": "run",
        }

    if kind == "table":
        registry = registry or Registry()
        table_id = spec.get("table_id")
        if not registry.has(table_id):
            raise UserError("数据表不存在或已删除。")
        source, table = registry.get(table_id)
        headers, rows, _ = registry.data(table_id)
        return {
            "name": registry.label(table_id),
            "headers": headers,
            "rows": rows,
            "kind": "table",
        }
    raise UserError("未知的对比对象。")


def _key_index(side: dict, keys: list[str]) -> list[int]:
    indices = []
    for key in keys:
        try:
            indices.append(side["headers"].index(key))
        except ValueError:
            raise UserError(f"「{side['name']}」中找不到主键列「{key}」。")
    return indices


def _row_map(rows: list[list[Any]], indices: list[int]) -> tuple[dict, int, int]:
    mapping = {}
    duplicate = 0
    skipped = 0
    for row in rows:
        key = tuple(norm_key(row[i] if i < len(row) else None) for i in indices)
        label = " / ".join(display_value(row[i] if i < len(row) else None) for i in indices)
        if not label.strip():
            skipped += 1
            continue
        if key in mapping:
            duplicate += 1
            continue
        mapping[key] = row
    return mapping, duplicate, skipped


def compare(left_spec: dict, right_spec: dict, keys: list[str]) -> dict:
    keys = [str(k).strip() for k in keys if str(k).strip()]
    if not keys:
        raise UserError("请至少选择一个主键列用于对比。")
    registry = Registry()
    left = load_side(left_spec, registry)
    right = load_side(right_spec, registry)
    left_keys = _key_index(left, keys)
    right_keys = _key_index(right, keys)

    left_common = [h for h in left["headers"]]
    right_common = [h for h in right["headers"]]
    only_left_cols = [h for h in left_common if h not in right_common]
    only_right_cols = [h for h in right_common if h not in left_common]
    common_cols = [h for h in left_common if h in right_common]
    compare_cols = [h for h in common_cols if h not in keys]

    lmap, ldup, lskip = _row_map(left["rows"], left_keys)
    rmap, rdup, rskip = _row_map(right["rows"], right_keys)
    lindex = {h: i for i, h in enumerate(left["headers"])}
    rindex = {h: i for i, h in enumerate(right["headers"])}

    added_keys = [k for k in rmap if k not in lmap]
    removed_keys = [k for k in lmap if k not in rmap]
    common_keys = [k for k in lmap if k in rmap]

    def key_label(key_tuple):
        row = lmap.get(key_tuple) or rmap.get(key_tuple)
        if not row:
            return " / ".join(key_tuple)
        mapping = lmap if key_tuple in lmap else rmap
        indices = left_keys if mapping is lmap else right_keys
        return " / ".join(display_value(row[i] if i < len(row) else None) for i in indices)

    all_changes = []
    changed_rows = 0
    changed_cells = 0
    unchanged = 0
    for key in common_keys:
        lrow, rrow = lmap[key], rmap[key]
        row_changed = False
        for col in compare_cols:
            lv = lrow[lindex[col]] if lindex[col] < len(lrow) else None
            rv = rrow[rindex[col]] if rindex[col] < len(rrow) else None
            if norm_key(lv) != norm_key(rv):
                row_changed = True
                changed_cells += 1
                all_changes.append({
                    "kind": "changed",
                    "key": key_label(key),
                    "field": col,
                    "left": display_value(lv),
                    "right": display_value(rv),
                })
        if row_changed:
            changed_rows += 1
        else:
            unchanged += 1

    added_rows = []
    for key in added_keys[:MAX_ROW_PREVIEW]:
        row = rmap[key]
        values = {}
        for h in right["headers"][:30]:
            values[h] = display_value(row[rindex[h]] if rindex[h] < len(row) else None)
        added_rows.append({"key": key_label(key), "values": values})
    removed_rows = []
    for key in removed_keys[:MAX_ROW_PREVIEW]:
        row = lmap[key]
        values = {}
        for h in left["headers"][:30]:
            values[h] = display_value(row[lindex[h]] if lindex[h] < len(row) else None)
        removed_rows.append({"key": key_label(key), "values": values})

    summary = {
        "left_rows": len(left["rows"]),
        "right_rows": len(right["rows"]),
        "added": len(added_keys),
        "removed": len(removed_keys),
        "changed_rows": changed_rows,
        "changed_cells": changed_cells,
        "unchanged": unchanged,
        "duplicate_left": ldup,
        "duplicate_right": rdup,
        "skipped_left": lskip,
        "skipped_right": rskip,
        "common_columns": len(common_cols),
        "only_left_columns": only_left_cols,
        "only_right_columns": only_right_cols,
        "keys": keys,
        "compared_at": now_iso(),
        "left_name": left["name"],
        "right_name": right["name"],
    }
    warnings = []
    if ldup or rdup:
        warnings.append(f"存在重复主键（左 {ldup} 个 / 右 {rdup} 个），对比时各取第一条。")
    if lskip or rskip:
        warnings.append(f"有 {lskip + rskip} 行主键为空，已跳过。")
    if only_left_cols or only_right_cols:
        warnings.append("两侧列不完全一致，仅对比共同列；单侧列已单独列出。")

    return {
        "summary": summary,
        "changes": all_changes[:MAX_CHANGE_ROWS],
        "added": added_rows,
        "removed": removed_rows,
        "warnings": warnings,
        "truncated": {
            "changes": len(all_changes) > MAX_CHANGE_ROWS,
            "added": len(added_keys) > MAX_ROW_PREVIEW,
            "removed": len(removed_keys) > MAX_ROW_PREVIEW,
        },
        "full": {
            "left": left,
            "right": right,
            "left_keys": left_keys,
            "right_keys": right_keys,
            "added_keys": added_keys,
            "removed_keys": removed_keys,
            "changes": all_changes,
            "has_more_changes": len(all_changes) > MAX_CHANGE_ROWS,
            "common_cols": common_cols,
            "only_left_cols": only_left_cols,
            "only_right_cols": only_right_cols,
        },
    }


def _style_header(ws, row: int, count: int):
    for c in range(1, count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = 26


def export_compare(result: dict) -> dict:
    summary = result["summary"]
    full = result["full"]
    left, right = full["left"], full["right"]
    compare_id = uid("cmp_")
    COMPARE_DIR.mkdir(parents=True, exist_ok=True)
    path = COMPARE_DIR / f"对比_{compare_id}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "对比摘要"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 66
    ws["A1"] = "表格汇总助手 · 数据对比报告"
    ws["A1"].font = Font(name="微软雅黑", size=15, bold=True, color=BRAND)
    ws.merge_cells("A1:B1")
    rows = [
        ("左侧数据", summary["left_name"]),
        ("右侧数据", summary["right_name"]),
        ("对比主键", " / ".join(summary["keys"])),
        ("对比时间", summary["compared_at"].replace("T", " ")[:19]),
        ("左 / 右行数", f"{summary['left_rows']} / {summary['right_rows']}"),
        ("新增行", str(summary["added"])),
        ("删除行", str(summary["removed"])),
        ("修改行", str(summary["changed_rows"])),
        ("修改单元格", str(summary["changed_cells"])),
        ("未变化行", str(summary["unchanged"])),
        ("仅左侧列", "、".join(summary["only_left_columns"]) or "—"),
        ("仅右侧列", "、".join(summary["only_right_columns"]) or "—"),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(name="微软雅黑", size=10, bold=True, color="55635D")
        ws.cell(row=i, column=2, value=v).font = BODY_FONT
        ws.cell(row=i, column=2).alignment = Alignment(vertical="center", wrap_text=True)

    detail = wb.create_sheet("差异明细")
    detail.sheet_view.showGridLines = False
    detail_headers = ["类型", "主键", "字段", "左侧值", "右侧值"]
    for i, h in enumerate(detail_headers, start=1):
        detail.cell(row=1, column=i, value=h)
    _style_header(detail, 1, len(detail_headers))
    detail_widths = [10, 24, 18, 30, 30]
    for i, w in enumerate(detail_widths, start=1):
        detail.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for change in full["changes"]:
        values = ["修改", change["key"], change["field"], change["left"], change["right"]]
        for c, value in enumerate(values, start=1):
            cell = detail.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.fill = FILL_CHG
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        r += 1
    for item in result["added"]:
        text = "；".join(f"{k}: {v}" for k, v in item["values"].items() if v)
        values = ["新增", item["key"], "整行", "—", text]
        for c, value in enumerate(values, start=1):
            cell = detail.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.fill = FILL_ADD
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1
    for item in result["removed"]:
        text = "；".join(f"{k}: {v}" for k, v in item["values"].items() if v)
        values = ["删除", item["key"], "整行", text, "—"]
        for c, value in enumerate(values, start=1):
            cell = detail.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.fill = FILL_DEL
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:E{max(1, r - 1)}"

    # 新增 / 删除完整行，便于逐行核对
    for title, side, keys_subset, fill, key_indices in (
        ("新增行", right, full["added_keys"], FILL_ADD, full["right_keys"]),
        ("删除行", left, full["removed_keys"], FILL_DEL, full["left_keys"]),
    ):
        sheet = wb.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        headers = side["headers"]
        for i, h in enumerate(headers, start=1):
            sheet.cell(row=1, column=i, value=h)
        _style_header(sheet, 1, len(headers))
        index = {h: i for i, h in enumerate(headers)}
        row_by_key = {}
        for original in side["rows"]:
            k = tuple(norm_key(original[i] if i < len(original) else None)
                      for i in key_indices)
            row_by_key.setdefault(k, original)
        for r_i, key in enumerate(keys_subset, start=2):
            original = row_by_key.get(key)
            if not original:
                continue
            for c, h in enumerate(headers, start=1):
                cell = sheet.cell(row=r_i, column=c, value=display_value(original[index[h]]))
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.fill = fill
        sheet.freeze_panes = "A2"
        for i, h in enumerate(headers, start=1):
            width = max(10, min(30, len(str(h)) * 2 + 4))
            sheet.column_dimensions[get_column_letter(i)].width = width

    wb.save(path)
    name = f"对比报告_{dt.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    db.execute(
        "INSERT INTO compares(id, name, path, created_at, summary) VALUES(?, ?, ?, ?, ?)",
        (compare_id, name, str(path), now_iso(), db.dumps(summary)),
    )
    return {"path": str(path), "name": name, "id": compare_id}
