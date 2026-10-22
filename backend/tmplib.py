"""数据模板：字段契约、校验规则、版本与引用关系。

## 与「汇总方案」的分工（这条界线是模块化的关键）

模板只回答**目标表长什么样**：有哪些列、每列什么类型、必填与否、有什么校验
规则、列宽与日期格式、以及可选的 Excel 样式底稿。它不关心数据从哪来，也不做
关联、合并、去重——那些是汇总方案的事。

方案单向依赖模板：方案的 `target.template_id` 指向模板，模板提供字段契约，
方案负责把数据源映射到这些字段上。所以：

- 删除被引用的模板会被拦住（除非调用方显式 force）；
- 模板字段改名/增删后，方案里会出现「字段与模板不一致」，可一键同步；
- 校验规则在生成/预览时执行，问题以警告呈现，不阻断生成。

## 为什么字段定义放在 meta.fields 而不是单独建表

模板的字段是**整体使用**的（一次编辑就是一份完整结构），单独建表会让「读取
一个模板」变成 N+1 次查询，而收益只有「按字段搜索」这一项——列表页的搜索
直接遍历 meta 里的字段名即可，量级完全够用。
"""
from __future__ import annotations

import datetime as dt
import re
import shutil
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import db
from .config import TEMPLATE_DIR
from .templates import (
    MAX_FIELD_NAME,
    MAX_TEMPLATE_FIELDS,
    delete_template_dir,
    fields_from_sheet,
    template_by_id,
)
from .utils import UserError, now_iso, parse_date_value, safe_filename, uid
from .writer import DATE_FORMATS

# ---- 可选项词汇表（前后端共用一份，避免两边对不上） ----------------------
FIELD_TYPES: dict[str, str] = {
    "auto": "自动识别",
    "text": "文本",
    "number": "数值",
    "date": "日期",
    "bool": "是/否",
}

RULE_KINDS: dict[str, str] = {
    "unique": "不允许重复",
    "enum": "只允许这些值",
    "pattern": "正则匹配",
    "min": "数值下限",
    "max": "数值上限",
    "min_len": "最短字符数",
    "max_len": "最长字符数",
}

TEMPLATE_STATUS = {"active": "使用中", "draft": "草稿", "archived": "已归档"}

MAX_TAGS = 8
MAX_TAG_LEN = 16
MAX_VERSIONS = 20
MAX_VALIDATE_ROWS = 20000
MAX_VIOLATIONS = 200

# 生成表格时的观感，与汇总产物保持一致
BRAND = "0F6B5C"
LINE = "DDE6E1"
INK = "1D2A25"
THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10.5, color=INK)
HEADER_ALIGN = Alignment(horizontal="left", vertical="center")

_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


# =========================================================================
# 字段定义：清洗与校验
# =========================================================================
def _clean_tags(raw: Any) -> list[str]:
    if raw is None:
        return []
    if isinstance(raw, str):
        raw = re.split(r"[,，;；\s]+", raw)
    if not isinstance(raw, (list, tuple)):
        raise UserError("标签格式不正确，应为数组或逗号分隔的文本。")
    out: list[str] = []
    for item in raw:
        tag = str(item or "").strip().lstrip("#").strip()
        if not tag:
            continue
        if len(tag) > MAX_TAG_LEN:
            raise UserError(f"标签「{tag[:8]}…」太长，单个标签请控制在 {MAX_TAG_LEN} 字以内。")
        if tag not in out:
            out.append(tag)
    if len(out) > MAX_TAGS:
        raise UserError(f"标签最多 {MAX_TAGS} 个。")
    return out


def _clean_rules(field_name: str, raw: Any) -> list[dict]:
    if raw in (None, ""):
        return []
    if not isinstance(raw, (list, tuple)):
        raise UserError(f"字段「{field_name}」的校验规则格式不正确。")
    rules: list[dict] = []
    seen: set[str] = set()
    for item in raw:
        if isinstance(item, str):
            item = {"kind": item}
        if not isinstance(item, dict):
            continue
        kind = str(item.get("kind") or "").strip()
        if kind not in RULE_KINDS:
            raise UserError(f"字段「{field_name}」有无法识别的校验规则：{kind or '（空）'}")
        if kind in seen:
            continue
        seen.add(kind)
        value = item.get("value")
        if value is None:
            value = item.get("values")  # 手写 JSON 里常写成 values，宽容一点
        message = str(item.get("message") or "").strip()[:80]
        if kind == "enum":
            if isinstance(value, str):
                value = re.split(r"[,，;；\n]+", value)
            if not isinstance(value, (list, tuple)):
                raise UserError(f"字段「{field_name}」的可选值格式不正确。")
            options = [str(v).strip() for v in value if str(v).strip()]
            if not options:
                raise UserError(f"字段「{field_name}」设置了「只允许这些值」，但没有填可选值。")
            if len(options) > 200:
                raise UserError(f"字段「{field_name}」的可选值最多 200 个。")
            rules.append({"kind": kind, "value": options, "message": message})
        elif kind == "pattern":
            pattern = str(value or "")
            if not pattern:
                raise UserError(f"字段「{field_name}」的正则规则没有填表达式。")
            try:
                re.compile(pattern)
            except re.error as exc:
                raise UserError(f"字段「{field_name}」的正则表达式有误：{exc}")
            rules.append({"kind": kind, "value": pattern, "message": message})
        elif kind in ("min", "max"):
            try:
                number = float(value)
            except (TypeError, ValueError):
                raise UserError(f"字段「{field_name}」的数值界限必须是数字。")
            rules.append({"kind": kind, "value": number, "message": message})
        elif kind in ("min_len", "max_len"):
            try:
                length = int(value)
            except (TypeError, ValueError):
                raise UserError(f"字段「{field_name}」的长度限制必须是整数。")
            if length < 0:
                raise UserError(f"字段「{field_name}」的长度限制不能为负数。")
            rules.append({"kind": kind, "value": length, "message": message})
        else:
            rules.append({"kind": kind, "value": True, "message": message})
    return rules


def normalize_fields(raw: Any, existing: list[dict] | None = None) -> list[dict]:
    """把前端传来的字段数组清洗成可信结构，顺便挡住坏数据。"""
    if raw is None:
        return []
    if not isinstance(raw, (list, tuple)):
        raise UserError("字段定义格式不正确，应为数组。")
    if len(raw) > MAX_TEMPLATE_FIELDS:
        raise UserError(f"字段太多，单个模板最多 {MAX_TEMPLATE_FIELDS} 个。")
    old_by_key = {f.get("key"): f for f in (existing or []) if isinstance(f, dict)}
    out: list[dict] = []
    seen: set[str] = set()
    for index, item in enumerate(raw):
        if not isinstance(item, dict):
            raise UserError(f"第 {index + 1} 个字段格式不正确。")
        name = str(item.get("name") or "").strip()
        if not name:
            raise UserError(f"第 {index + 1} 个字段没有填列名。")
        if len(name) > MAX_FIELD_NAME:
            raise UserError(f"列名「{name[:10]}…」太长，请控制在 {MAX_FIELD_NAME} 字以内。")
        if name in seen:
            raise UserError(f"列名「{name}」重复了，同一张表里列名必须唯一。")
        seen.add(name)
        type_ = str(item.get("type") or "auto")
        if type_ not in FIELD_TYPES:
            type_ = "auto"
        try:
            width = int(item.get("width") or 14)
        except (TypeError, ValueError):
            width = 14
        width = max(6, min(60, width))
        date_format = str(item.get("date_format") or "")
        if date_format and date_format not in DATE_FORMATS:
            date_format = ""
        key = str(item.get("key") or "").strip() or uid("f_")
        if key in old_by_key and old_by_key[key].get("key"):
            key = old_by_key[key]["key"]
        col = item.get("template_col")
        try:
            col = int(col) if col is not None and str(col) != "" else None
        except (TypeError, ValueError):
            col = None
        out.append({
            "key": key,
            "name": name,
            "type": type_,
            "required": bool(item.get("required")),
            "default": str(item.get("default") or "")[:200],
            "width": width,
            "date_format": date_format,
            "note": str(item.get("note") or item.get("description") or "")[:200],
            "sample": str(item.get("sample") or "")[:120],
            "rules": _clean_rules(name, item.get("rules")),
            "template_col": col if col is not None else index,
        })
    return out


def fields_signature(fields: list[dict]) -> list[tuple]:
    """只比较会影响数据写入的部分，避免改个说明就产生一个版本。"""
    return [
        (f.get("name"), f.get("type"), bool(f.get("required")), str(f.get("default") or ""),
         f.get("width"), f.get("date_format"), f.get("template_col"),
         db.dumps(f.get("rules") or []))
        for f in fields
    ]


def rules_count(fields: list[dict]) -> int:
    return sum(len(f.get("rules") or []) + (1 if f.get("required") else 0) for f in fields)


# =========================================================================
# 由字段定义生成 Excel 模板
# =========================================================================
def _safe_sheet_name(name: str) -> str:
    cleaned = _SHEET_BAD.sub(" ", str(name or "")).strip() or "汇总"
    return cleaned[:31]


def build_workbook(fields: list[dict], sheet_name: str, path: Path, title: str = "") -> dict:
    """按字段定义生成一份可用的 .xlsx 模板。

    生成物刻意保持「干净」：第 1 行表头、第 2 行起是数据区，不加说明行——
    说明行会被 writer 当成第一条数据，反而添乱。字段说明放在第二个工作表。
    """
    sheet_name = _safe_sheet_name(sheet_name)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.showGridLines = False

    for i, field in enumerate(fields, start=1):
        cell = ws.cell(row=1, column=i, value=field["name"])
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(i)].width = field.get("width") or 14
        fmt = DATE_FORMATS.get(field.get("date_format") or "")
        if fmt:
            for row in range(2, 2001):
                ws.cell(row=row, column=i).number_format = fmt
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    # 枚举与布尔字段挂下拉，省得录入时手滑
    for i, field in enumerate(fields, start=1):
        options: list[str] = []
        for rule in field.get("rules") or []:
            if rule.get("kind") == "enum":
                options = [str(v) for v in (rule.get("value") or [])]
        if field.get("type") == "bool":
            options = ["是", "否"]
        if not options:
            continue
        joined = ",".join(o.replace(",", "，") for o in options)[:250]
        dv = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=not field.get("required"))
        dv.error = "请从下拉列表中选择"
        dv.errorTitle = "值不在允许范围内"
        ws.add_data_validation(dv)
        col = get_column_letter(i)
        dv.add(f"{col}2:{col}2000")

    info = wb.create_sheet("字段说明")
    info.sheet_view.showGridLines = False
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 10
    info.column_dimensions["C"].width = 8
    info.column_dimensions["D"].width = 46
    info.column_dimensions["E"].width = 26
    info["A1"] = f"{title or '数据模板'} · 字段说明"
    info["A1"].font = Font(name="微软雅黑", size=13, bold=True, color=BRAND)
    info.merge_cells("A1:E1")
    heads = ["字段", "类型", "必填", "校验规则", "说明 / 示例"]
    for i, head in enumerate(heads, start=1):
        cell = info.cell(row=3, column=i, value=head)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    for r, field in enumerate(fields, start=4):
        info.cell(row=r, column=1, value=field["name"]).font = BODY_FONT
        info.cell(row=r, column=2, value=FIELD_TYPES.get(field.get("type") or "auto")).font = BODY_FONT
        info.cell(row=r, column=3, value="是" if field.get("required") else "").font = BODY_FONT
        rules_text = "；".join(describe_rule(rule) for rule in field.get("rules") or [])
        info.cell(row=r, column=4, value=rules_text).font = BODY_FONT
        note = field.get("note") or ""
        sample = field.get("sample") or ""
        text = note + (f"（示例：{sample}）" if sample else "")
        info.cell(row=r, column=5, value=text).font = BODY_FONT

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {
        "name": sheet_name,
        "header_row": 1,
        "headers": [f["name"] for f in fields],
        "row_count": 0,
        "col_count": len(fields),
        "hidden": False,
    }


def describe_rule(rule: dict) -> str:
    kind = rule.get("kind")
    value = rule.get("value")
    if kind == "unique":
        return "不允许重复"
    if kind == "enum":
        return "仅限：" + " / ".join(str(v) for v in (value or [])[:6])
    if kind == "pattern":
        return f"匹配 {value}"
    if kind == "min":
        return f"≥ {value}"
    if kind == "max":
        return f"≤ {value}"
    if kind == "min_len":
        return f"至少 {value} 字"
    if kind == "max_len":
        return f"最多 {value} 字"
    return RULE_KINDS.get(kind, str(kind))


# =========================================================================
# 读取与引用关系
# =========================================================================
def references_map() -> dict[str, list[dict]]:
    """哪些汇总方案引用了哪些模板。单一来源，列表页与详情页共用。"""
    out: dict[str, list[dict]] = {}
    for row in db.query("SELECT id, name, config, updated_at FROM profiles"):
        config = db.loads(row.get("config"), {}) or {}
        target = config.get("target") or {}
        template_id = target.get("template_id")
        if not template_id:
            continue
        out.setdefault(template_id, []).append({
            "id": row["id"],
            "name": row["name"],
            "mode": config.get("mode") or "join",
            "sheet": target.get("sheet") or "",
            "updated_at": row.get("updated_at"),
        })
    return out


def _decorate(row: dict, refs: list[dict]) -> dict:
    meta = row.get("meta") or {}
    fields = meta.get("fields") or []
    sheets = meta.get("sheets") or []
    tags = row.get("tags") or []
    if isinstance(tags, str):
        tags = db.loads(tags, []) or []
    item = {
        "id": row["id"],
        "name": row["name"],
        "description": row.get("description") or "",
        "category": row.get("category") or "",
        "tags": tags,
        "status": row.get("status") or "active",
        "origin": row.get("origin") or "upload",
        "version": int(row.get("version") or 1),
        "size": row.get("size") or 0,
        "created_at": row.get("created_at"),
        "updated_at": row.get("updated_at") or row.get("created_at"),
        "sheet": meta.get("sheet") or (sheets[0]["name"] if sheets else ""),
        "header_row": meta.get("header_row") or (sheets[0].get("header_row") if sheets else 1) or 1,
        "sheet_count": len(sheets),
        "sheets": [
            {k: s.get(k) for k in ("name", "header_row", "headers", "row_count", "col_count")}
            for s in sheets
        ],
        "fields": fields,
        "field_count": len(fields),
        "rules_count": rules_count(fields),
        "has_file": bool(row.get("path")),
        "ref_count": len(refs),
        "refs": refs[:4],
    }
    # 兼容旧结构：早期前端与测试都走 template.meta.sheets / meta.fields。
    # 由上面的规范字段派生，不额外存库，避免两份数据各自演化。
    item["meta"] = {
        "sheets": item["sheets"],
        "fields": fields,
        "sheet": item["sheet"],
        "header_row": item["header_row"],
        "warnings": meta.get("warnings") or [],
    }
    return item


def list_templates(q: str = "", category: str = "", tag: str = "", status: str = "all",
                   sort: str = "updated") -> dict:
    refs = references_map()
    items = [_decorate(row, refs.get(row["id"], [])) for row in db.template_rows()]

    categories: dict[str, int] = {}
    tags: dict[str, int] = {}
    for item in items:
        if item["category"]:
            categories[item["category"]] = categories.get(item["category"], 0) + 1
        for t in item["tags"]:
            tags[t] = tags.get(t, 0) + 1

    keyword = (q or "").strip().lower()
    if keyword:
        def hit(item: dict) -> bool:
            haystack = " ".join([
                item["name"], item["description"], item["category"],
                " ".join(item["tags"]),
                " ".join(f.get("name") or "" for f in item["fields"]),
            ]).lower()
            return keyword in haystack
        items = [i for i in items if hit(i)]
    if category:
        items = [i for i in items if i["category"] == category]
    if tag:
        items = [i for i in items if tag in i["tags"]]
    if status and status != "all":
        items = [i for i in items if i["status"] == status]

    if sort == "name":
        items.sort(key=lambda i: i["name"])
    elif sort == "fields":
        items.sort(key=lambda i: i["field_count"], reverse=True)
    elif sort == "refs":
        items.sort(key=lambda i: i["ref_count"], reverse=True)
    elif sort == "created":
        items.sort(key=lambda i: i["created_at"] or "", reverse=True)
    else:
        items.sort(key=lambda i: i["updated_at"] or "", reverse=True)

    return {
        "templates": items,
        "total": len(items),
        "categories": [{"name": k, "count": v} for k, v in sorted(categories.items(), key=lambda x: -x[1])],
        "tags": [{"name": k, "count": v} for k, v in sorted(tags.items(), key=lambda x: -x[1])],
        "stats": {
            "total": len(items),
            "active": sum(1 for i in items if i["status"] == "active"),
            "fields": sum(i["field_count"] for i in items),
            "referenced": sum(1 for i in items if i["ref_count"]),
        },
    }


def template_detail(template_id: str) -> dict:
    row = template_by_id(template_id)
    if not row:
        raise UserError("模板不存在或已删除。")
    refs = references_map().get(template_id, [])
    item = _decorate(row, refs)
    item["file_name"] = Path(row.get("path") or "").name
    item["versions"] = [
        {**v, "snapshot": _snapshot_brief(v.get("snapshot") or {})}
        for v in db.template_version_rows(template_id)
    ]
    return item


def public_template(template_id: str) -> dict:
    """不含版本历史的对外形态，供上传/删除等接口返回。"""
    row = template_by_id(template_id)
    if not row:
        raise UserError("模板不存在或已删除。")
    return _decorate(row, references_map().get(template_id, []))


def _snapshot_brief(snapshot: dict) -> dict:
    fields = snapshot.get("fields") or []
    return {
        "name": snapshot.get("name"),
        "category": snapshot.get("category"),
        "tags": snapshot.get("tags") or [],
        "sheet": snapshot.get("sheet"),
        "field_count": len(fields),
        "rules_count": rules_count(fields),
        "fields": [{"name": f.get("name"), "type": f.get("type")} for f in fields[:30]],
    }


def _snapshot(row: dict) -> dict:
    meta = row.get("meta") or {}
    return {
        "name": row.get("name"),
        "description": row.get("description") or "",
        "category": row.get("category") or "",
        "tags": row.get("tags") or [],
        "status": row.get("status") or "active",
        "sheet": meta.get("sheet") or "",
        "header_row": meta.get("header_row") or 1,
        "fields": meta.get("fields") or [],
    }


# =========================================================================
# 版本
# =========================================================================
def _push_version(template_id: str, row: dict, note: str) -> None:
    version = int(row.get("version") or 1)
    db.execute(
        "INSERT INTO template_versions(id, template_id, version, snapshot, note, created_at) "
        "VALUES(?, ?, ?, ?, ?, ?)",
        (uid("tv_"), template_id, version, db.dumps(_snapshot(row)), note[:120], now_iso()),
    )
    # 只留最近 N 个版本，避免历史无限膨胀
    stale = db.query(
        "SELECT id FROM template_versions WHERE template_id=? "
        "ORDER BY version DESC, created_at DESC LIMIT -1 OFFSET ?",
        (template_id, MAX_VERSIONS),
    )
    for item in stale:
        db.execute("DELETE FROM template_versions WHERE id=?", (item["id"],))


def save_version(template_id: str, note: str = "") -> dict:
    row = template_by_id(template_id)
    if not row:
        raise UserError("模板不存在或已删除。")
    _push_version(template_id, row, note or "手动记录")
    return {"versions": template_detail(template_id)["versions"]}


def restore_version(template_id: str, version_id: str) -> dict:
    row = template_by_id(template_id)
    if not row:
        raise UserError("模板不存在或已删除。")
    record = db.query_one(
        "SELECT * FROM template_versions WHERE id=? AND template_id=?", (version_id, template_id))
    if not record:
        raise UserError("找不到这个历史版本。")
    snapshot = db.loads(record.get("snapshot"), {}) or {}
    fields = normalize_fields(snapshot.get("fields") or [], row["meta"].get("fields"))
    if not fields:
        raise UserError("这个历史版本里没有字段定义，无法回滚。")
    _push_version(template_id, row, f"回滚到 v{record.get('version')} 之前")
    _apply_fields(row, fields, snapshot.get("sheet") or row["meta"].get("sheet") or "汇总",
                  int(snapshot.get("header_row") or 1), regenerate=True)
    db.execute(
        "UPDATE templates SET name=?, description=?, category=?, tags=?, version=?, updated_at=? WHERE id=?",
        (str(snapshot.get("name") or row["name"])[:80],
         str(snapshot.get("description") or "")[:300],
         str(snapshot.get("category") or "")[:40],
         db.dumps(snapshot.get("tags") or row.get("tags") or []),
         int(row.get("version") or 1) + 1,
         now_iso(), template_id),
    )
    return template_detail(template_id)


# =========================================================================
# 字段落盘：字段变了，文件也要跟着变
# =========================================================================
def _apply_fields(row: dict, fields: list[dict], sheet: str, header_row: int,
                  regenerate: bool, rename_headers: bool = True) -> None:
    """把字段写回 meta，并按 origin 决定是否重画 Excel 文件。

    - origin='fields'：文件是程序生成的，字段变了就整份重画（没有用户样式可丢）。
    - origin='upload'：文件是用户传的，只改表头单元格的值，样式一律不动。
    - rename_headers=False：这次变化来自「重新读取表头」，文件才是真相，
      不能反过来把字段名写回文件，否则会串列。
    """
    meta = dict(row.get("meta") or {})
    old_fields = meta.get("fields") or []
    meta["fields"] = fields
    meta["sheet"] = sheet
    meta["header_row"] = header_row
    sheets = meta.get("sheets") or []
    target = next((s for s in sheets if s.get("name") == sheet), None)
    path = Path(row.get("path") or "")

    if row.get("origin") == "fields" or regenerate or not path.exists():
        build_workbook(fields, sheet, path, title=row.get("name") or "")
        meta["sheets"] = [{
            "name": _safe_sheet_name(sheet),
            "header_row": header_row,
            "headers": [f["name"] for f in fields],
            "row_count": 0,
            "col_count": len(fields),
            "hidden": False,
        }]
    else:
        renamed = [
            (old.get("template_col"), old.get("name"), new.get("name"))
            for old, new in zip(old_fields, fields)
            if old.get("name") != new.get("name")
        ]
        if renamed and rename_headers:
            _rename_headers(path, sheet, header_row, renamed)
        if target is not None:
            target["headers"] = [f["name"] for f in fields]
            target["col_count"] = len(fields)
        elif sheets:
            sheets[0]["headers"] = [f["name"] for f in fields]
        meta["sheets"] = sheets
    db.execute("UPDATE templates SET meta=? WHERE id=?", (db.dumps(meta), row["id"]))


def _rename_headers(path: Path, sheet: str, header_row: int, renamed: list[tuple]) -> None:
    keep_vba = path.suffix.lower() == ".xlsm"
    try:
        wb = load_workbook(path, keep_vba=keep_vba)
    except Exception as exc:
        raise UserError(f"无法打开模板文件，字段名改动未保存：{exc}") from exc
    if sheet not in wb.sheetnames:
        raise UserError("模板文件里找不到该工作表，字段名改动未保存。")
    ws = wb[sheet]
    for col_index, _old, new in renamed:
        try:
            index = int(col_index) + 1
        except (TypeError, ValueError):
            continue
        if index < 1:
            continue
        cell = ws.cell(row=header_row, column=index)
        try:
            cell.value = new
        except AttributeError:
            raise UserError("模板表头存在合并单元格，请先在 Excel 里取消合并再改名。")
    wb.save(path)


# =========================================================================
# 增删改
# =========================================================================
def create_template(payload: dict) -> dict:
    name = str(payload.get("name") or "").strip()
    if not name:
        raise UserError("请填写模板名称。")
    fields = normalize_fields(payload.get("fields") or [])
    if not fields:
        raise UserError("请至少定义一个字段。")
    sheet = _safe_sheet_name(payload.get("sheet") or "汇总")
    template_id = uid("t_")
    directory = TEMPLATE_DIR / template_id
    directory.mkdir(parents=True, exist_ok=True)
    filename = safe_filename(f"{name[:40]}.xlsx")
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    path = directory / filename
    sheet_meta = build_workbook(fields, sheet, path, title=name)
    ts = now_iso()
    meta = {"sheets": [sheet_meta], "fields": fields, "sheet": sheet_meta["name"],
            "header_row": 1, "warnings": []}
    db.execute(
        "INSERT INTO templates(id, name, path, size, created_at, updated_at, description, "
        "category, tags, status, origin, version, meta) "
        "VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'fields', 1, ?)",
        (template_id, name[:80], str(path), path.stat().st_size, ts, ts,
         str(payload.get("description") or "")[:300],
         str(payload.get("category") or "")[:40],
         db.dumps(_clean_tags(payload.get("tags"))),
         "active" if payload.get("status") not in TEMPLATE_STATUS else payload.get("status"),
         db.dumps(meta)),
    )
    return template_detail(template_id)


def update_template(template_id: str, payload: dict) -> dict:
    row = template_by_id(template_id)
    if not row:
        raise UserError("模板不存在或已删除。")
    meta = row.get("meta") or {}
    old_fields = meta.get("fields") or []

    new_fields = old_fields
    structural = False
    if payload.get("fields") is not None:
        new_fields = normalize_fields(payload.get("fields"), old_fields)
        if not new_fields:
            raise UserError("模板至少要保留一个字段。")
        structural = fields_signature(new_fields) != fields_signature(old_fields)

    # 从文件重新读取表头（用户调整了表头行或换了工作表）
    resynced = False
    if payload.get("resync"):
        from .templates import inspect_sheet
        sheet_name = payload.get("sheet") or meta.get("sheet") or "汇总"
        header_row = payload.get("header_row") or meta.get("header_row") or 1
        info = inspect_sheet(template_id, sheet_name, int(header_row), old_fields)
        new_fields = normalize_fields(info.get("fields") or [], old_fields)
        structural = True
        resynced = True
        meta["sheet"] = info["name"]
        meta["header_row"] = info.get("header_row") or header_row

    if structural:
        _push_version(template_id, row, str(payload.get("note") or "字段结构变更"))

    name = str(payload.get("name") or row["name"]).strip()[:80] or row["name"]
    description = payload.get("description", row.get("description") or "")
    category = payload.get("category", row.get("category") or "")
    tags = payload.get("tags", row.get("tags") or [])
    status = payload.get("status", row.get("status") or "active")
    if status not in TEMPLATE_STATUS:
        status = "active"

    db.execute(
        "UPDATE templates SET name=?, description=?, category=?, tags=?, status=?, updated_at=? WHERE id=?",
        (name, str(description or "")[:300], str(category or "")[:40],
         db.dumps(_clean_tags(tags)), status, now_iso(), template_id),
    )
    row = template_by_id(template_id)
    fields_touched = payload.get("fields") is not None or resynced
    if fields_touched:
        sheet = meta.get("sheet") or row["meta"].get("sheet") or "汇总"
        header_row = int(meta.get("header_row") or row["meta"].get("header_row") or 1)
        if structural:
            _apply_fields(row, new_fields, sheet, header_row,
                          regenerate=(row.get("origin") == "fields"),
                          rename_headers=not resynced)
            db.execute("UPDATE templates SET version=? WHERE id=?",
                       (int(row.get("version") or 1) + 1, template_id))
        else:
            # 只改了说明/示例值这类不影响写入的内容：不重画文件、不加版本号，
            # 但字段本身必须落库，否则用户改完一刷新就白改了。
            _store_fields_only(row, new_fields, sheet, header_row)
    return template_detail(template_id)


def _store_fields_only(row: dict, fields: list[dict], sheet: str, header_row: int) -> None:
    meta = dict(row.get("meta") or {})
    meta["fields"] = fields
    meta["sheet"] = sheet
    meta["header_row"] = header_row
    db.execute("UPDATE templates SET meta=? WHERE id=?", (db.dumps(meta), row["id"]))


def duplicate_template(template_id: str, name: str = "") -> dict:
    row = template_by_id(template_id)
    if not row:
        raise UserError("模板不存在或已删除。")
    new_id = uid("t_")
    directory = TEMPLATE_DIR / new_id
    directory.mkdir(parents=True, exist_ok=True)
    src = Path(row.get("path") or "")
    new_name = (name or f"{row['name']} 副本")[:80]
    suffix = src.suffix or ".xlsx"
    target = directory / safe_filename(f"{new_name[:40]}{suffix}")
    if src.exists():
        shutil.copy2(src, target)
    ts = now_iso()
    db.execute(
        "INSERT INTO templates(id, name, path, size, created_at, updated_at, description, "
        "category, tags, status, origin, version, meta) "
        "VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?)",
        (new_id, new_name, str(target), target.stat().st_size if target.exists() else 0,
         ts, ts, row.get("description") or "", row.get("category") or "",
         db.dumps(row.get("tags") or []), row.get("status") or "active",
         row.get("origin") or "upload", db.dumps(row.get("meta") or {})),
    )
    return template_detail(new_id)


def delete_template(template_id: str, force: bool = False) -> dict:
    row = template_by_id(template_id)
    if not row:
        raise UserError("模板不存在或已删除。")
    refs = references_map().get(template_id, [])
    if refs and not force:
        names = "、".join(r["name"] for r in refs[:3])
        more = f" 等 {len(refs)} 个" if len(refs) > 3 else ""
        raise UserError(
            f"还有汇总方案在用这个模板（{names}{more}）。请先在方案里换掉模板，"
            "或确认后选择「仍然删除」。"
        )
    delete_template_dir(row)
    db.execute("DELETE FROM template_versions WHERE template_id=?", (template_id,))
    db.execute("DELETE FROM templates WHERE id=?", (template_id,))
    return {"deleted": template_id, "forced": bool(force and refs)}


# =========================================================================
# 导入 / 导出
# =========================================================================
EXPORT_SCHEMA = 1


def export_template(template_id: str) -> dict:
    detail = template_detail(template_id)
    return {
        "schema": EXPORT_SCHEMA,
        "kind": "huiyi-template",
        "exported_at": now_iso(),
        "template": {
            "name": detail["name"],
            "description": detail["description"],
            "category": detail["category"],
            "tags": detail["tags"],
            "sheet": detail["sheet"],
            "header_row": detail["header_row"],
            "fields": detail["fields"],
        },
    }


def import_template(payload: dict, name_prefix: str = "") -> dict:
    if not isinstance(payload, dict):
        raise UserError("导入文件格式不正确。")
    if payload.get("kind") not in (None, "huiyi-template"):
        raise UserError("这不是表格汇总助手导出的模板文件。")
    if int(payload.get("schema") or 1) > EXPORT_SCHEMA:
        raise UserError("这份模板来自更新版本的程序，请先升级后再导入。")
    data = payload.get("template") or payload
    name = str(data.get("name") or "").strip()
    if not name:
        raise UserError("导入内容缺少模板名称。")
    fields = data.get("fields") or []
    if not fields:
        raise UserError("导入内容缺少字段定义。")
    existing = {t["name"] for t in db.template_rows()}
    final_name = f"{name_prefix}{name}"[:80]
    if final_name in existing:
        index = 2
        while f"{final_name} ({index})" in existing:
            index += 1
        final_name = f"{final_name} ({index})"
    return create_template({
        "name": final_name,
        "description": data.get("description") or "",
        "category": data.get("category") or "",
        "tags": data.get("tags") or [],
        "sheet": data.get("sheet") or "汇总",
        "fields": fields,
    })


# =========================================================================
# 字段校验：生成/预览时对数据行执行
# =========================================================================
def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def check_value(field: dict, value: Any) -> tuple[bool, str]:
    """单值校验。返回 (是否合规, 人话原因)。"""
    if _is_blank(value):
        if field.get("required"):
            return False, "必填项为空"
        return True, ""
    text = str(value).strip()
    type_ = field.get("type") or "auto"
    if type_ == "number":
        try:
            float(str(value).replace(",", ""))
        except (TypeError, ValueError):
            return False, f"不是数字：{text[:20]}"
    elif type_ == "date":
        if not _looks_like_date(value):
            return False, f"不是日期：{text[:20]}"
    elif type_ == "bool":
        if text not in ("是", "否", "Y", "N", "y", "n", "true", "false", "1", "0", "True", "False"):
            return False, f"不是是/否值：{text[:20]}"
    for rule in field.get("rules") or []:
        kind = rule.get("kind")
        target = rule.get("value")
        if kind == "enum":
            options = [str(v) for v in (target or [])]
            if text not in options:
                return False, rule.get("message") or f"不在允许值内：{text[:20]}"
        elif kind == "pattern":
            try:
                if not re.fullmatch(str(target), text):
                    return False, rule.get("message") or f"格式不符：{text[:20]}"
            except re.error:
                continue
        elif kind in ("min", "max"):
            try:
                number = float(str(value).replace(",", ""))
            except (TypeError, ValueError):
                continue
            if kind == "min" and number < float(target):
                return False, rule.get("message") or f"小于下限 {target}"
            if kind == "max" and number > float(target):
                return False, rule.get("message") or f"大于上限 {target}"
        elif kind == "min_len" and len(text) < int(target):
            return False, rule.get("message") or f"长度不足 {target}"
        elif kind == "max_len" and len(text) > int(target):
            return False, rule.get("message") or f"长度超过 {target}"
    return True, ""


def _looks_like_date(value: Any) -> bool:
    """日期类型校验直接借用写入层的解析器，保证「校验通过」与「写进 Excel」
    是同一套判断——否则 2026-02-30 这种会在校验里蒙混过关、写入时变成文本。
    """
    if isinstance(value, (dt.date, dt.datetime)):
        return True
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel 日期序列值（约 2009 ~ 2064），模板里可能已是数字格式的日期
        return 40000 <= float(value) <= 60000
    ok, _parsed = parse_date_value(value)
    return ok


def validate_rows(fields: list[dict], headers: list[str], rows: list[list],
                  first_data_row: int = 2, max_rows: int = MAX_VALIDATE_ROWS) -> dict:
    """按模板字段规则检查输出数据。

    约定：`headers` 是输出表头（顺序与 rows 的列一一对应），`first_data_row`
    是第一条数据在 Excel 里的行号，报错信息里带上它，用户才能直接定位。
    """
    if not fields:
        return {"checked_rows": 0, "fields_checked": 0, "count": 0,
                "violations": [], "by_field": [], "truncated": False, "skipped": "模板没有字段定义"}
    index = {str(h).strip(): i for i, h in enumerate(headers)}
    checks = [(f, index[f["name"]]) for f in fields if f.get("name") in index]
    if not checks:
        return {"checked_rows": 0, "fields_checked": 0, "count": 0,
                "violations": [], "by_field": [], "truncated": False,
                "skipped": "输出列与模板字段没有交集"}

    violations: list[dict] = []
    by_field: dict[str, int] = {}
    seen_unique: dict[str, dict[str, int]] = {f["name"]: {} for f in fields if _has_kind(f, "unique")}
    scanned = rows[:max_rows]
    total = 0
    for offset, row in enumerate(scanned):
        excel_row = first_data_row + offset
        for field, col in checks:
            value = row[col] if col < len(row) else None
            ok, reason = check_value(field, value)
            if ok and value is not None and field["name"] in seen_unique:
                key = str(value).strip()
                if key:
                    first = seen_unique[field["name"]].get(key)
                    if first is not None:
                        ok, reason = False, f"与第 {first} 行重复：{key[:20]}"
                    else:
                        seen_unique[field["name"]][key] = excel_row
            if ok:
                continue
            total += 1
            by_field[field["name"]] = by_field.get(field["name"], 0) + 1
            if len(violations) < MAX_VIOLATIONS:
                violations.append({
                    "row": excel_row,
                    "field": field["name"],
                    "reason": reason,
                    "value": "" if value is None else str(value)[:40],
                })
    return {
        "checked_rows": len(scanned),
        "fields_checked": len(checks),
        "count": total,
        "violations": violations,
        "by_field": sorted(
            ({"field": k, "count": v} for k, v in by_field.items()),
            key=lambda x: -x["count"],
        ),
        "truncated": total > len(violations) or len(rows) > len(scanned),
        "unique_fields": list(seen_unique),
    }


def _has_kind(field: dict, kind: str) -> bool:
    return any(r.get("kind") == kind for r in (field.get("rules") or []))


def validate_config_rows(config: dict, result: dict) -> dict | None:
    """给预览/运行用的薄封装：只有目标表是模板时才校验。"""
    target = config.get("target") or {}
    if target.get("kind") != "template":
        return None
    row = template_by_id(target.get("template_id") or "")
    if not row:
        return None
    fields = (row.get("meta") or {}).get("fields") or []
    rows = result.get("rows") or result.get("preview_rows") or []
    header_row = int(target.get("header_row") or 1)
    return validate_rows(fields, result.get("headers") or [], rows, first_data_row=header_row + 1)


__all__ = [
    "FIELD_TYPES", "RULE_KINDS", "TEMPLATE_STATUS", "EXPORT_SCHEMA",
    "build_workbook", "create_template", "delete_template", "describe_rule",
    "duplicate_template", "export_template", "fields_from_sheet", "fields_signature",
    "import_template", "list_templates", "normalize_fields",
    "references_map", "restore_version", "rules_count", "save_version",
    "template_by_id", "template_detail", "update_template", "validate_config_rows",
    "validate_rows", "check_value", "DATE_FORMATS",
]
