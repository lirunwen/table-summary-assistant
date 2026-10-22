"""数据模板模块：CRUD、版本、引用保护、导入导出、校验规则。

这组用例盯的是「模板 = 目标表结构契约」这条边界：
- 字段定义能存能改能回滚，且改完之后 Excel 表头真的跟着变；
- 被方案引用时删不掉，除非显式 force；
- 校验规则只提示不阻断，且日期用和写入时同一套解析。
"""
from __future__ import annotations

import json

from openpyxl import load_workbook

from backend import db, tmplib
from backend.config import TEMPLATE_DIR


def post(client, url, payload=None, method="post"):
    resp = getattr(client, method)(url, json=payload or {})
    assert resp.status_code == 200, (url, resp.status_code, resp.get_json())
    data = resp.get_json()
    assert data.get("ok"), data
    return data


def make_template(client, name="人员信息表", fields=None, **extra):
    payload = {
        "name": name,
        "fields": fields or [
            {"name": "员工编号", "type": "text", "required": True,
             "rules": [{"kind": "unique"}, {"kind": "pattern", "value": r"^WX\d{6,7}$"}]},
            {"name": "姓名", "type": "text", "required": True},
            {"name": "所属部门", "type": "text"},
            {"name": "入职日期", "type": "date"},
            {"name": "月薪", "type": "number", "width": 12,
             "rules": [{"kind": "min", "value": 0}, {"kind": "max", "value": 200000}]},
        ],
        **extra,
    }
    return post(client, "/api/templates", payload)["template"]


# --------------------------------------------------------------------------
# 创建与读取
# --------------------------------------------------------------------------
def test_create_template_generates_xlsx_and_contract(client):
    tpl = make_template(client, category="人事", tags=["月度", "常用"],
                        description="月度花名册")
    assert tpl["origin"] == "fields"
    assert tpl["version"] == 1
    assert tpl["field_count"] == 5
    assert tpl["sheet"] == "汇总"
    assert tpl["category"] == "人事"
    assert tpl["tags"] == ["月度", "常用"]
    # 规则数 = 必填 2 + unique 1 + pattern 1 + min 1 + max 1
    assert tpl["rules_count"] == 6

    path = TEMPLATE_DIR / tpl["id"] / tpl["file_name"]
    assert path.exists()
    wb = load_workbook(path)
    assert "汇总" in wb.sheetnames
    assert "字段说明" in wb.sheetnames          # 第二张表给使用方看
    ws = wb["汇总"]
    headers = [c.value for c in ws[1]]
    assert headers == ["员工编号", "姓名", "所属部门", "入职日期", "月薪"]
    assert ws.freeze_panes == "A2"


def test_enum_field_gets_excel_dropdown(client):
    tpl = make_template(client, fields=[
        {"name": "职级", "type": "text", "required": True,
         "rules": [{"kind": "enum", "values": ["13A", "12B", "14A"]}]},
    ])
    wb = load_workbook(TEMPLATE_DIR / tpl["id"] / tpl["file_name"])
    validations = wb["汇总"].data_validations.dataValidation
    assert len(validations) == 1
    assert validations[0].type == "list"
    assert "13A" in (validations[0].formula1 or "")


def test_create_rejects_bad_input(client):
    cases = [
        ({"name": "", "fields": [{"name": "A"}]}, "模板名称"),
        ({"name": "X", "fields": []}, "字段"),
        ({"name": "X", "fields": [{"name": ""}]}, "列名"),
        ({"name": "X", "fields": [{"name": "A"}, {"name": "A"}]}, "重复"),
        ({"name": "X", "fields": [{"name": "A", "rules": [{"kind": "pattern", "value": "("}]}]}, "正则"),
        ({"name": "X", "fields": [{"name": "A", "rules": [{"kind": "enum", "value": []}]}]}, "可选值"),
        ({"name": "X", "fields": [{"name": "A", "rules": [{"kind": "min", "value": "abc"}]}]}, "数字"),
        ({"name": "X", "fields": [{"name": "A", "rules": [{"kind": "nope"}]}]}, "无法识别"),
    ]
    for payload, keyword in cases:
        resp = client.post("/api/templates", json=payload)
        assert resp.status_code == 400, payload
        assert keyword in resp.get_json()["error"], (payload, resp.get_json())


def test_max_field_name_and_width_are_clamped(client):
    tpl = make_template(client, fields=[
        {"name": "A", "width": 999},
        {"name": "B", "width": 1},
    ])
    widths = [f["width"] for f in tpl["fields"]]
    assert widths == [60, 6]


# --------------------------------------------------------------------------
# 列表：搜索 / 筛选 / 排序
# --------------------------------------------------------------------------
def test_list_search_hits_field_names(client):
    make_template(client, name="人员信息表", category="人事", tags=["月度"])
    make_template(client, name="计件工资表", category="财务", tags=["计件"],
                  fields=[{"name": "件数"}, {"name": "单价"}])
    # 「部门」只出现在第一个模板的字段里，不在名称/说明里
    hit = client.get("/api/templates?q=部门").get_json()
    assert [t["name"] for t in hit["templates"]] == ["人员信息表"]

    by_cat = client.get("/api/templates?category=财务").get_json()
    assert [t["name"] for t in by_cat["templates"]] == ["计件工资表"]

    by_tag = client.get("/api/templates?tag=月度").get_json()
    assert [t["name"] for t in by_tag["templates"]] == ["人员信息表"]

    assert hit["categories"] and any(c["name"] == "人事" for c in hit["categories"])
    assert any(t["name"] == "月度" for t in hit["tags"])


def test_list_sort_and_status_filter(client):
    a = make_template(client, name="A 表")
    b = make_template(client, name="B 表", fields=[{"name": f"c{i}"} for i in range(1, 8)])
    post(client, f"/api/templates/{a['id']}", {"status": "archived"}, "put")

    archived = client.get("/api/templates?status=archived").get_json()
    assert [t["name"] for t in archived["templates"]] == ["A 表"]

    by_fields = client.get("/api/templates?sort=fields").get_json()
    assert [t["name"] for t in by_fields["templates"]][0] == "B 表"

    by_name = client.get("/api/templates?sort=name").get_json()
    assert [t["name"] for t in by_name["templates"]] == ["A 表", "B 表"]
    assert b["id"] != a["id"]


# --------------------------------------------------------------------------
# 编辑：改字段 → 表头跟着改 + 自动存版本
# --------------------------------------------------------------------------
def test_rename_field_rewrites_header_and_snapshots(client):
    tpl = make_template(client)
    path = TEMPLATE_DIR / tpl["id"] / tpl["file_name"]

    fields = [dict(f) for f in tpl["fields"]]
    fields[2]["name"] = "服务组织"
    updated = post(client, f"/api/templates/{tpl['id']}", {"fields": fields}, "put")["template"]

    assert updated["version"] == 2
    assert len(updated["versions"]) == 1
    assert updated["versions"][0]["snapshot"]["field_count"] == 5

    assert tpl["id"] in str(path)
    headers = [c.value for c in load_workbook(path)["汇总"][1]]
    assert headers[2] == "服务组织"


def test_cosmetic_edit_does_not_bump_version(client):
    tpl = make_template(client)
    fields = [dict(f) for f in tpl["fields"]]
    fields[0]["note"] = "只是加个说明"
    fields[0]["sample"] = "WX9071526"
    updated = post(client, f"/api/templates/{tpl['id']}", {"fields": fields}, "put")["template"]
    assert updated["version"] == 1
    assert updated["versions"] == []
    assert updated["fields"][0]["note"] == "只是加个说明"


def test_restore_version_rolls_back_fields(client):
    tpl = make_template(client)
    fields = [dict(f) for f in tpl["fields"]]
    fields.append({"name": "备注", "type": "text"})
    v2 = post(client, f"/api/templates/{tpl['id']}", {"fields": fields}, "put")["template"]
    assert v2["field_count"] == 6

    version_id = v2["versions"][0]["id"]
    restored = post(client, f"/api/templates/{tpl['id']}/versions/{version_id}/restore", {})["template"]
    assert restored["field_count"] == 5
    assert "备注" not in [f["name"] for f in restored["fields"]]
    # 回滚本身也是一次结构变化，旧结构会再存一份快照
    assert len(restored["versions"]) >= 2


def test_manual_snapshot(client):
    tpl = make_template(client)
    data = post(client, f"/api/templates/{tpl['id']}/versions", {"note": "上线前"})["template"]
    assert data["versions"][0]["note"] == "上线前"
    detail = post(client, f"/api/templates/{tpl['id']}", {"description": "改说明"}, "put")["template"]
    assert len(detail["versions"]) == 1
    assert detail["description"] == "改说明"


# --------------------------------------------------------------------------
# 复制 / 删除 / 引用保护
# --------------------------------------------------------------------------
def test_duplicate_keeps_fields_and_suffixes_name(client):
    tpl = make_template(client, name="人员信息表")
    dup = post(client, f"/api/templates/{tpl['id']}/duplicate", {})["template"]
    assert dup["id"] != tpl["id"]
    assert dup["name"] == "人员信息表 副本"
    assert [f["name"] for f in dup["fields"]] == [f["name"] for f in tpl["fields"]]


def test_delete_is_blocked_while_referenced(client):
    tpl = make_template(client)
    cfg = {
        "mode": "join", "primary_table": "", "primary_keys": [], "lookups": [],
        "append_tables": [], "auto_by_name": True,
        "target": {"kind": "template", "template_id": tpl["id"], "sheet": "汇总",
                   "header_row": 1, "write_mode": "replace", "columns": []},
        "mappings": {}, "options": {"dedupe": True, "skip_empty_key": False},
    }
    post(client, "/api/profiles", {"name": "月度花名册", "config": cfg})

    got = client.get(f"/api/templates/{tpl['id']}").get_json()["template"]
    assert got["ref_count"] == 1
    assert got["refs"][0]["name"] == "月度花名册"

    resp = client.delete(f"/api/templates/{tpl['id']}")
    assert resp.status_code == 400
    assert "月度花名册" in resp.get_json()["error"]

    resp = client.delete(f"/api/templates/{tpl['id']}?force=1")
    assert resp.status_code == 200
    assert tpl["id"] not in [t["id"] for t in resp.get_json()["templates"]]
    gone = client.get(f"/api/templates/{tpl['id']}")
    assert gone.status_code == 400
    assert "不存在" in gone.get_json()["error"]
    # 文件和目录一起清掉
    assert not (TEMPLATE_DIR / tpl["id"]).exists()


# --------------------------------------------------------------------------
# 导出 / 导入
# --------------------------------------------------------------------------
def test_export_then_import_roundtrip(client):
    tpl = make_template(client, name="人员信息表", category="人事", tags=["月度"])
    resp = client.get(f"/api/templates/{tpl['id']}/export")
    assert resp.status_code == 200
    payload = json.loads(resp.data.decode("utf-8"))
    assert payload["kind"] == "huiyi-template"
    assert payload["schema"] == 1
    assert [f["name"] for f in payload["template"]["fields"]] == [f["name"] for f in tpl["fields"]]

    imported = post(client, "/api/templates/import", payload)["template"]
    # 同名自动加后缀，不会覆盖原件
    assert imported["name"] == "人员信息表 (2)"
    assert imported["id"] != tpl["id"]
    assert [f["name"] for f in imported["fields"]] == [f["name"] for f in tpl["fields"]]


def test_import_rejects_foreign_json(client):
    resp = client.post("/api/templates/import", json={"hello": "world"})
    assert resp.status_code == 400
    assert "模板" in resp.get_json()["error"]


# --------------------------------------------------------------------------
# 上传 Excel 模板：origin=upload，改字段只改表头单元格
# --------------------------------------------------------------------------
def test_upload_template_derives_fields_from_headers(client, sample):
    with open(sample("人员汇总模板.xlsx"), "rb") as fh:
        resp = client.post("/api/templates/upload",
                           data={"file": (fh, "人员汇总模板.xlsx"), "category": "人事"},
                           content_type="multipart/form-data")
    tpl = resp.get_json()["template"]
    assert tpl["origin"] == "upload"
    assert tpl["field_count"] > 0
    assert tpl["category"] == "人事"
    assert all(f["name"] for f in tpl["fields"])


def test_uploaded_template_keeps_user_rules_on_header_change(client, sample):
    with open(sample("人员汇总模板.xlsx"), "rb") as fh:
        tpl = client.post("/api/templates/upload",
                          data={"file": (fh, "人员汇总模板.xlsx")},
                          content_type="multipart/form-data").get_json()["template"]
    fields = [dict(f) for f in tpl["fields"]]
    fields[0] = {**fields[0], "required": True, "type": "text",
                 "rules": [{"kind": "unique"}]}
    saved = post(client, f"/api/templates/{tpl['id']}", {"fields": fields}, "put")["template"]
    assert saved["fields"][0]["required"] is True

    # 从文件重读表头：字段配置不该被清空
    resynced = post(client, f"/api/templates/{tpl['id']}",
                    {"sheet": saved["sheet"], "resync": True}, "put")["template"]
    assert resynced["fields"][0]["required"] is True
    assert resynced["fields"][0]["rules"][0]["kind"] == "unique"


def test_resync_does_not_write_header_back(client, sample):
    """resync 是「以文件为准」，不能反手把字段名写回文件导致串列。"""
    with open(sample("人员汇总模板.xlsx"), "rb") as fh:
        tpl = client.post("/api/templates/upload",
                          data={"file": (fh, "人员汇总模板.xlsx")},
                          content_type="multipart/form-data").get_json()["template"]
    path = TEMPLATE_DIR / tpl["id"] / tpl["file_name"]

    fields = [dict(f) for f in tpl["fields"]]
    fields[0]["name"] = "手工改过的名字"
    post(client, f"/api/templates/{tpl['id']}", {"fields": fields}, "put")
    before = [c.value for c in load_workbook(path)[tpl["sheet"]][tpl["header_row"]]]
    assert before[0] == "手工改过的名字"

    post(client, f"/api/templates/{tpl['id']}", {"sheet": tpl["sheet"], "resync": True}, "put")

    after = [c.value for c in load_workbook(path)[tpl["sheet"]][tpl["header_row"]]]
    assert after == before


# --------------------------------------------------------------------------
# 校验：只提示不阻断
# --------------------------------------------------------------------------
def test_validate_reports_violations_without_blocking(client):
    tpl = make_template(client)
    payload = {
        "rows": [
            ["WX9071526", "张三", "GTC", "2026-03-01", 26500],
            ["WX9071526", "", "GTC", "2026-02-30", "abc"],   # 重复 / 缺姓名 / 非法日期 / 非数字
            ["ABC", "李四", "GTC", "2026-05-01", 30000],      # 工号格式不对
        ],
    }
    payload = {"headers": [f["name"] for f in tpl["fields"]], **payload}
    data = post(client, f"/api/templates/{tpl['id']}/validate", payload)
    assert data["ok"] is True                     # 有违规也是 200
    result = data["validation"]
    assert result["checked_rows"] == 3
    assert result["count"] >= 4
    fields = {v["field"] for v in result["violations"]}
    assert {"员工编号", "姓名", "入职日期", "月薪"} <= fields
    assert any("重复" in v["reason"] for v in result["violations"])


def test_validate_empty_value_only_triggers_required(client):
    tpl = make_template(client, fields=[
        {"name": "工号", "type": "text", "required": True,
         "rules": [{"kind": "pattern", "value": r"^WX\d+$"}, {"kind": "min_len", "value": 8}]},
    ])
    data = post(client, f"/api/templates/{tpl['id']}/validate",
                {"headers": ["工号"], "rows": [[""]]})["validation"]
    assert data["count"] == 1
    assert "必填" in data["violations"][0]["reason"]


def test_date_validation_matches_writer_rules():
    """校验用的日期判断必须和写 Excel 时同一套，否则会出现「校验过了但写不进去」。"""
    tpl_field = {"name": "日期", "type": "date", "rules": []}
    ok, _ = tmplib.check_value(tpl_field, "2026-02-30")
    assert ok is False                     # 2 月没有 30 号
    ok, _ = tmplib.check_value(tpl_field, "2026-02-28")
    assert ok is True
    ok, _ = tmplib.check_value(tpl_field, 45000)   # Excel 日期序列号
    assert ok is True
    ok, _ = tmplib.check_value(tpl_field, "下周三")
    assert ok is False


# --------------------------------------------------------------------------
# 与方案联动：preview / run 的 summary 带 field_validation
# --------------------------------------------------------------------------
def test_validate_config_rows_only_runs_for_template_target(client):
    tpl = make_template(client)
    result = {"headers": [f["name"] for f in tpl["fields"]],
              "rows": [["WX123456", "张三", "GTC", "2026-01-01", 100]]}
    ok_cfg = {"target": {"kind": "template", "template_id": tpl["id"], "header_row": 1}}
    report = tmplib.validate_config_rows(ok_cfg, result)
    assert report["checked_rows"] == 1
    assert report["count"] == 0

    # 目标表是空白表 / 模板已删时不做校验，避免无谓计算
    assert tmplib.validate_config_rows({"target": {"kind": "blank"}}, result) is None
    assert tmplib.validate_config_rows(
        {"target": {"kind": "template", "template_id": "t_gone"}}, result) is None


# --------------------------------------------------------------------------
# 迁移：老库升级后模板表要有新列
# --------------------------------------------------------------------------
def test_migration_adds_template_columns(app):
    cols = {row["name"] for row in db.query("PRAGMA table_info(templates)")}
    assert {"description", "category", "tags", "status", "origin",
            "version", "updated_at"} <= cols
    tables = {row["name"] for row in db.query(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    assert "template_versions" in tables
