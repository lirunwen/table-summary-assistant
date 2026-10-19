"""后端端到端冒烟测试：上传、模板、关联汇总、对比、导出。"""
from __future__ import annotations

import json
import os
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "_smoke"
shutil.rmtree(DATA, ignore_errors=True)
os.environ["HUIYI_DATA_DIR"] = str(DATA)
sys.path.insert(0, str(ROOT))

from backend import db  # noqa: E402
from backend.api import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

db.init_db()
app = create_app(port=0)
app.config.update(TESTING=True)
client = app.test_client()

samples = ROOT / "samples"


def post_json(url, payload, method="post"):
    resp = getattr(client, method)(url, json=payload)
    data = resp.get_json()
    assert resp.status_code == 200, (url, resp.status_code, data)
    assert data.get("ok"), (url, data)
    return data


def upload_sources():
    files = []
    for name in ["员工主表.xlsx", "考勤统计.xlsx", "工资表.csv", "通讯录更新.eml"]:
        files.append((open(samples / name, "rb"), name))
    resp = client.post("/api/sources/upload", data={"files": files},
                       content_type="multipart/form-data")
    data = resp.get_json()
    assert resp.status_code == 200, data
    assert data["ok"], data
    print("上传:", [(s["name"], s["kind"], len(s["meta"]["tables"])) for s in data["created"]])
    assert not data["errors"], data["errors"]
    return data


data = upload_sources()
boot = client.get("/api/bootstrap").get_json()
assert boot["ok"]
for s in boot["sources"]:
    print("源:", s["name"], s["kind"], [(t["label"], t["header_row"], t["headers"], t["row_count"]) for t in s["meta"]["tables"]])
    for c in s.get("children") or []:
        print(" 附件:", c["name"], [(t["label"], t["headers"], t["row_count"]) for t in c["meta"]["tables"]])

email = next(s for s in boot["sources"] if s["kind"] == "email")
assert len(email["meta"]["tables"]) == 1, email["meta"]["tables"]
assert len(email.get("children") or []) == 1
assert email["meta"]["tables"][0]["row_count"] == 5

# 模板
with open(samples / "人员汇总模板.xlsx", "rb") as fh:
    resp = client.post("/api/templates/upload", data={"file": (fh, "人员汇总模板.xlsx")},
                       content_type="multipart/form-data")
tpl_data = resp.get_json()
assert tpl_data["ok"], tpl_data
tpl = tpl_data["template"]
print("模板:", tpl["name"], tpl["meta"]["sheets"])
sheet = tpl["meta"]["sheets"][0]
assert sheet["header_row"] == 2
assert sheet["headers"][0] == "工号"

sources = boot["sources"]
emp = next(s for s in sources if s["name"] == "员工主表.xlsx")
att = next(s for s in sources if s["name"] == "考勤统计.xlsx")
sal = next(s for s in sources if s["name"] == "工资表.csv")
emp_t = emp["meta"]["tables"][0]
att_t = att["meta"]["tables"][0]
sal_t = sal["meta"]["tables"][0]
mail_t = email["meta"]["tables"][0]
att_child = email["children"][0]
child_t = att_child["meta"]["tables"][0]

assert emp_t["header_row"] == 2, emp_t
assert att_t["header_row"] == 3, att_t
assert sal_t["headers"][0] == "工号", sal_t

columns = []
for i, h in enumerate(sheet["headers"]):
    columns.append({
        "key": f"t{i}", "name": h, "type": "auto", "width": 14,
        "included": True, "template_col": i,
    })

config = {
    "mode": "join",
    "primary_table": emp_t["id"],
    "primary_keys": [0],
    "lookups": [
        {"table_id": att_t["id"], "keys": [0]},
        {"table_id": sal_t["id"], "keys": [0]},
        {"table_id": mail_t["id"], "keys": [0]},
    ],
    "append_tables": [],
    "auto_by_name": True,
    "target": {
        "kind": "template", "template_id": tpl["id"], "sheet": sheet["name"],
        "header_row": 2, "write_mode": "replace", "columns": columns,
    },
    "mappings": {
        "t0": {"mode": "table", "table_id": emp_t["id"], "column_index": 0, "transform": "trim", "fallback": ""},
        "t1": {"mode": "table", "table_id": emp_t["id"], "column_index": 1, "transform": "direct", "fallback": ""},
        "t2": {"mode": "table", "table_id": emp_t["id"], "column_index": 2, "transform": "direct", "fallback": ""},
        "t3": {"mode": "table", "table_id": emp_t["id"], "column_index": 3, "transform": "direct", "fallback": ""},
        "t4": {"mode": "table", "table_id": att_t["id"], "column_index": 1, "transform": "number", "fallback": ""},
        "t5": {"mode": "table", "table_id": att_t["id"], "column_index": 2, "transform": "number", "fallback": ""},
        "t6": {"mode": "table", "table_id": sal_t["id"], "column_index": 1, "transform": "number", "fallback": ""},
        "t7": {"mode": "const", "value": "2026-10"},
    },
    "options": {"dedupe": True, "skip_empty_key": True},
}

profile = post_json("/api/profiles", {"name": "10月人员汇总", "mode": "join"})["profile"]
config["name"] = profile["name"]
profile = post_json(f"/api/profiles/{profile['id']}", {"name": profile["name"], "config": config}, method="put")["profile"]

preview = post_json("/api/preview", {"config": config, "limit": 5})
print("预览:", preview["stats"], preview["headers"])
assert preview["stats"]["rows"] == 18, preview["stats"]
assert len(preview["rows"]) == 5
assert preview["rows"][0][0] == "E001"

run = post_json(f"/api/profiles/{profile['id']}/run", {"config": config})["run"]
print("运行:", run["summary"], run["output_name"])
assert run["summary"]["rows"] == 18
assert run["summary"]["duplicates"] == 0
summary = run["summary"]
assert summary["output"]["header_row"] == 2
assert summary["output"]["columns"] == list(range(1, 9))

# 下载并检查输出
resp = client.get(run["download_url"])
assert resp.status_code == 200
out_path = DATA / "out.xlsx"
out_path.write_bytes(resp.data)
wb = load_workbook(out_path, data_only=True)
ws = wb["人员汇总"]
assert ws["A2"].value == "工号"
assert ws["A3"].value == "E001", ws["A3"].value
assert ws["A20"].value == "E018", ws["A20"].value
assert ws["A3"].fill.fgColor.rgb.endswith("0F6B5C") is False
assert ws.max_row == 20, ws.max_row
assert ws["H3"].value == "2026-10"
assert ws["G20"].value in (None, ""), ws["G20"].value
assert ws["E3"].value == 21
print("输出:", ws.max_row, "行；E001 =", [ws.cell(row=3, column=c).value for c in range(1, 9)])

# 表头调整接口
resp = client.patch(f"/api/sources/{att['id']}", json={
    "table_id": att_t["id"], "header_row": 3,
})
assert resp.get_json()["ok"], resp.get_json()
print("表头调整 OK")

# 再次运行一个备注不同的版本，用于对比
config2 = json.loads(json.dumps(config))
config2["mappings"]["t7"] = {"mode": "const", "value": "2026-11"}
run2 = post_json(f"/api/profiles/{profile['id']}/run", {"config": config2})["run"]
assert run2["summary"]["rows"] == 18

compare = post_json("/api/compare", {
    "left": {"type": "run", "id": run["id"]},
    "right": {"type": "run", "id": run2["id"]},
    "keys": ["工号"],
})
print("对比:", compare["summary"])
assert compare["summary"]["changed_rows"] == 18
assert compare["summary"]["changed_cells"] == 18
assert compare["summary"]["added"] == 0
resp = client.get(compare["export_url"])
assert resp.status_code == 200
cmp_path = DATA / "compare.xlsx"
cmp_path.write_bytes(resp.data)
wb2 = load_workbook(cmp_path, data_only=True)
assert "对比摘要" in wb2.sheetnames
assert "差异明细" in wb2.sheetnames
print("对比报告:", wb2.sheetnames, wb2["差异明细"].max_row)

# 对比自身
same = post_json("/api/compare", {
    "left": {"type": "run", "id": run["id"]},
    "right": {"type": "run", "id": run["id"]},
    "keys": ["工号"],
})
assert same["summary"]["changed_cells"] == 0
print("自身对比 OK")

# 纵向合并
config3 = {
    "mode": "append",
    "primary_table": "",
    "primary_keys": [],
    "lookups": [],
    "append_tables": [emp_t["id"], att_t["id"]],
    "auto_by_name": False,
    "target": {
        "kind": "blank", "template_id": None, "sheet": "合并结果",
        "header_row": 1, "write_mode": "replace",
        "columns": [
            {"key": "a0", "name": "工号", "type": "auto", "width": 12, "included": True},
            {"key": "a1", "name": "姓名", "type": "auto", "width": 12, "included": True},
        ],
    },
    "mappings": {
        "a0": {"mode": "table", "table_id": emp_t["id"], "column_index": 0, "transform": "trim", "fallback": ""},
        "a1": {"mode": "table", "table_id": emp_t["id"], "column_index": 1, "transform": "direct", "fallback": ""},
    },
    "options": {"dedupe": False},
}
p3 = post_json("/api/profiles", {"name": "合并测试", "mode": "append"})["profile"]
run3 = post_json(f"/api/profiles/{p3['id']}/run", {"config": config3})["run"]
print("合并:", run3["summary"]["rows"])
assert run3["summary"]["rows"] >= 18

print("\n全部冒烟测试通过 ✅")

# ---------- 额外覆盖：结果 vs 数据表、纵向自动匹配、xls、msg、公式注入、模板追加 ----------
print("\n--- 额外测试 ---")

# 1) 结果表与原始数据表对比
cmp2 = post_json("/api/compare", {
    "left": {"type": "run", "id": run["id"]},
    "right": {"type": "table", "table_id": emp_t["id"]},
    "keys": ["工号"],
})
assert cmp2["summary"]["left_rows"] == 18
assert cmp2["summary"]["right_rows"] == 18
assert cmp2["summary"]["changed_cells"] == 0, cmp2["summary"]
assert "手机号" in cmp2["summary"]["only_right_columns"]
print("结果 vs 数据表对比 OK", cmp2["summary"]["only_left_columns"], cmp2["summary"]["only_right_columns"])

# 2) 纵向合并 + 同名列自动匹配
cfg_append = {
    "mode": "append", "primary_table": "", "primary_keys": [], "lookups": [],
    "append_tables": [emp_t["id"], sal_t["id"]], "auto_by_name": True,
    "target": {"kind": "blank", "template_id": None, "sheet": "合并结果",
               "header_row": 1, "write_mode": "replace", "columns": [
        {"key": "a0", "name": "工号", "type": "auto", "width": 12, "included": True},
        {"key": "a1", "name": "姓名", "type": "auto", "width": 12, "included": True},
        {"key": "a2", "name": "应发工资", "type": "number", "width": 14, "included": True},
    ]},
    "mappings": {
        "a0": {"mode": "table", "table_id": emp_t["id"], "column_index": 0, "transform": "trim", "fallback": ""},
        "a1": {"mode": "table", "table_id": emp_t["id"], "column_index": 1, "transform": "direct", "fallback": ""},
        "a2": {"mode": "table", "table_id": sal_t["id"], "column_index": 1, "transform": "number", "fallback": ""},
    },
    "options": {"dedupe": False},
}
p4 = post_json("/api/profiles", {"name": "纵向自动匹配", "mode": "append"})["profile"]
prev4 = post_json("/api/preview", {"config": cfg_append, "limit": 50})
assert prev4["stats"]["rows"] == 35, prev4["stats"]
assert prev4["rows"][18][0] == "E001", prev4["rows"][18]
assert prev4["rows"][18][2] == 23800, prev4["rows"][18]
print("纵向自动匹配 OK", prev4["rows"][18])

# 3) 旧版 .xls
with open(samples / "旧格式.xls", "rb") as fh:
    resp = client.post("/api/sources/upload", data={"files": [(fh, "旧格式.xls")]},
                       content_type="multipart/form-data")
xls_data = resp.get_json()
assert xls_data["ok"], xls_data
xls_source = next(s for s in xls_data["sources"] if s["name"] == "旧格式.xls")
xls_table = xls_source["meta"]["tables"][0]
assert xls_table["headers"] == ["工号", "姓名", "部门"], xls_table
assert xls_table["row_count"] == 2
print("旧版 .xls OK")

# 4) 真实 .msg 解析（无表格时应给出提示）
msg_path = Path("/tmp/multi-to-to.msg")
if msg_path.exists():
    with open(msg_path, "rb") as fh:
        resp = client.post("/api/sources/upload", data={"files": [(fh, "multi-to-to.msg")]},
                           content_type="multipart/form-data")
    msg_data = resp.get_json()
    assert msg_data["ok"], msg_data
    msg_source = next(s for s in msg_data["sources"] if s["name"] == "multi-to-to.msg")
    assert msg_source["kind"] == "email"
    assert msg_source["meta"]["email"]["subject"].startswith("Test: multiple To")
    assert msg_source["meta"]["warnings"], msg_source["meta"]
    print("Outlook .msg OK", msg_source["meta"]["email"]["subject"])

# 5) 公式注入防护
formula_csv = DATA / "formula.csv"
formula_csv.write_text("工号,备注\nF001,=1+1\nF002,@SUM(A1)\n", encoding="utf-8")
with open(formula_csv, "rb") as fh:
    resp = client.post("/api/sources/upload", data={"files": [(fh, "formula.csv")]},
                       content_type="multipart/form-data")
fdata = resp.get_json()
f_source = next(s for s in fdata["sources"] if s["name"] == "formula.csv")
f_table = f_source["meta"]["tables"][0]
cfg_f = {
    "mode": "join", "primary_table": f_table["id"], "primary_keys": [0], "lookups": [],
    "append_tables": [], "auto_by_name": True,
    "target": {"kind": "blank", "template_id": None, "sheet": "公式防护",
               "header_row": 1, "write_mode": "replace", "columns": [
        {"key": "f0", "name": "工号", "type": "text", "width": 12, "included": True},
        {"key": "f1", "name": "备注", "type": "text", "width": 20, "included": True},
    ]},
    "mappings": {
        "f0": {"mode": "table", "table_id": f_table["id"], "column_index": 0, "transform": "trim", "fallback": ""},
        "f1": {"mode": "table", "table_id": f_table["id"], "column_index": 1, "transform": "direct", "fallback": ""},
    },
    "options": {},
}
pf = post_json("/api/profiles", {"name": "公式防护", "mode": "join"})["profile"]
rf = post_json(f"/api/profiles/{pf['id']}/run", {"config": cfg_f})["run"]
resp = client.get(rf["download_url"])
fpath = DATA / "formula_out.xlsx"
fpath.write_bytes(resp.data)
wb_f = load_workbook(fpath, data_only=False)
ws_f = wb_f["公式防护"]
assert ws_f["B2"].value == "=1+1", ws_f["B2"].value
assert ws_f["B2"].data_type == "s", ws_f["B2"].data_type
assert ws_f["B3"].value == "@SUM(A1)", ws_f["B3"].value
print("公式注入防护 OK", ws_f["B2"].value, ws_f["B2"].data_type)

# 6) 模板追加模式
cfg_tpl_append = json.loads(json.dumps(config))
cfg_tpl_append["target"]["write_mode"] = "append"
run_ta = post_json(f"/api/profiles/{profile['id']}/run", {"config": cfg_tpl_append})["run"]
resp = client.get(run_ta["download_url"])
tapath = DATA / "template_append.xlsx"
tapath.write_bytes(resp.data)
wb_ta = load_workbook(tapath, data_only=True)
ws_ta = wb_ta["人员汇总"]
assert ws_ta["A3"].value == "E000", ws_ta["A3"].value
assert ws_ta["A4"].value == "E001", ws_ta["A4"].value
assert ws_ta.max_row == 21, ws_ta.max_row
print("模板追加模式 OK", ws_ta.max_row, ws_ta["A3"].value, ws_ta["A4"].value)

print("\n额外测试全部通过 ✅")
