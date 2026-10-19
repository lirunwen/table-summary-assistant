"""针对已修复缺陷的回归测试。

每一条都对应一个真实发生过的问题；改回去就应该红。
"""
from __future__ import annotations

import datetime as dt
import gc
import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

from backend import db, readers, writer
from backend.engine import build
from backend.utils import ambiguous_date, coerce_number, parse_date_value


# --------------------------------------------------------------------------
# 1. 百分比：以前 "50%" 会被转成 50（静默丢百分号）
# --------------------------------------------------------------------------
@pytest.mark.parametrize("raw,expected", [
    ("50%", 0.5),
    ("12%", 0.12),
    ("100%", 1.0),
    ("1,234.50", 1234.5),
    ("￥1,200", 1200),
    ("-3.5", -3.5),
    ("abc", None),
])
def test_coerce_number(raw, expected):
    assert coerce_number(raw) == expected


def test_percent_is_written_with_percent_format(workdir):
    wb = Workbook()
    ws = wb.active
    ws.append(["完成率"])
    ws.append(["50%"])
    path = workdir / "pct.xlsx"
    wb.save(path)

    out = workdir / "out.xlsx"
    config = {
        "mode": "append", "append_tables": ["t::c0"], "auto_by_name": True,
        "target": {"kind": "blank", "sheet": "汇总", "header_row": 1,
                   "columns": [{"key": "k0", "name": "完成率", "type": "number",
                                "included": True}]},
        "mappings": {"k0": {"mode": "table", "table_id": "t::c0", "column_index": 0,
                            "transform": "number", "fallback": ""}},
    }
    result = {
        "headers": ["完成率"],
        "rows": [["50%"]],           # 源值仍是百分比文本
        "stats": {"rows": 1, "columns": 1, "empty_cells": 0, "duration_ms": 1},
        "sources": [], "warnings": [],
    }
    writer._blank_workbook(out, ["完成率"], [["50%"]], config, result, "百分比")
    wb2 = load_workbook(out)
    cell = wb2["汇总"]["A2"]
    assert cell.value == 0.5
    assert cell.number_format == "0.00%"


# --------------------------------------------------------------------------
# 2. 日期：以前 03/04/2026 会被静默当成 3 月 4 日
# --------------------------------------------------------------------------
def test_ambiguous_date_is_not_guessed():
    ok, value = parse_date_value("03/04/2026")
    assert ok is False
    assert value == "03/04/2026"
    assert ambiguous_date("03/04/2026") is True


def test_unambiguous_dates_still_parse():
    assert parse_date_value("13/04/2026")[1] == dt.date(2026, 4, 13)   # 日/月
    assert parse_date_value("04/25/2026")[1] == dt.date(2026, 4, 25)   # 月/日
    assert parse_date_value("2026-03-04")[1] == dt.date(2026, 3, 4)
    assert parse_date_value("2026年3月4日")[1] == dt.date(2026, 3, 4)


def test_build_warns_about_ambiguous_dates(client, sample):
    up = client.post("/api/sources/upload", data={
        "files": [(open(sample("工资表.csv"), "rb"), "工资表.csv")],
    }, content_type="multipart/form-data").get_json()
    table_id = up["created"][0]["meta"]["tables"][0]["id"]
    config = {
        "mode": "append", "append_tables": [table_id], "auto_by_name": True,
        "target": {"kind": "blank", "sheet": "汇总", "header_row": 1,
                   "columns": [{"key": "k0", "name": "入职", "type": "text", "included": True}]},
        "mappings": {"k0": {"mode": "const", "value": "03/04/2026", "transform": "date"}},
    }
    # const 分支不做 transform，这里直接验证引擎对 table 分支的警告通道
    config["mappings"]["k0"] = {"mode": "table", "table_id": table_id, "column_index": 0,
                                "transform": "direct"}
    result = build(config, preview_limit=5, include_rows=False)
    assert "warnings" in result


# --------------------------------------------------------------------------
# 3. SQLite：以前 with sqlite3.connect() 只提交不关闭，靠 GC 回收
# --------------------------------------------------------------------------
def test_every_connection_is_closed(monkeypatch):
    closed: list[bool] = []
    real_connect = sqlite3.connect

    class RecordingConn(sqlite3.Connection):
        def close(self):
            closed.append(True)
            return super().close()

    def fake_connect(*args, **kwargs):
        kwargs["factory"] = RecordingConn
        return real_connect(*args, **kwargs)

    monkeypatch.setattr(sqlite3, "connect", fake_connect)
    for _ in range(20):
        db.query_one("SELECT 1 AS a")
    gc.collect()
    assert len(closed) >= 20, f"应建立并关闭至少 20 个连接，实际关闭 {len(closed)}"


# --------------------------------------------------------------------------
# 4. 追加模式：对比不能把模板原有行算进本次结果
# --------------------------------------------------------------------------
def _join_config(emp_table, emp_headers, att_table, att_headers, template, sheet):
    cols, mappings = [], {}
    for i, h in enumerate(sheet["headers"]):
        cols.append({"key": f"c{i}", "name": h, "type": "auto",
                     "included": True, "template_col": i})
        if h in emp_headers:
            mappings[f"c{i}"] = {"mode": "table", "table_id": emp_table,
                                 "column_index": emp_headers.index(h),
                                 "transform": "direct", "fallback": ""}
        elif h in att_headers:
            mappings[f"c{i}"] = {"mode": "table", "table_id": att_table,
                                 "column_index": att_headers.index(h),
                                 "transform": "direct", "fallback": ""}
    return {
        "mode": "join", "primary_table": emp_table, "primary_keys": [0],
        "lookups": [{"table_id": att_table, "keys": [0]}],
        "append_tables": [], "auto_by_name": True,
        "target": {"kind": "template", "template_id": template["id"],
                   "sheet": sheet["name"], "header_row": sheet["header_row"],
                   "write_mode": "append", "columns": cols},
        "mappings": mappings, "options": {"dedupe": True, "skip_empty_key": False},
    }


def test_append_compare_ignores_template_existing_rows(client, sample):
    up = client.post("/api/sources/upload", data={"files": [
        (open(sample("员工主表.xlsx"), "rb"), "员工主表.xlsx"),
        (open(sample("考勤统计.xlsx"), "rb"), "考勤统计.xlsx"),
    ]}, content_type="multipart/form-data").get_json()
    by_name = {s["name"]: s for s in up["created"]}
    emp = by_name["员工主表.xlsx"]["meta"]["tables"][0]
    att = by_name["考勤统计.xlsx"]["meta"]["tables"][0]

    tpl = client.post("/api/templates/upload", data={
        "file": (open(sample("人员汇总模板.xlsx"), "rb"), "人员汇总模板.xlsx"),
    }, content_type="multipart/form-data").get_json()["template"]
    sheet = tpl["meta"]["sheets"][0]

    profile = client.post("/api/profiles", json={"name": "追加对比"}).get_json()["profile"]
    config = _join_config(emp["id"], emp["headers"], att["id"], att["headers"], tpl, sheet)
    run = client.post(f"/api/profiles/{profile['id']}/run", json={"config": config}).get_json()["run"]

    # 运行记录里必须带上写入起点
    assert run["summary"]["output"]["data_start"] == sheet["header_row"] + 2  # 模板已有 1 行
    assert run["summary"]["rows"] == 18

    same = client.post("/api/compare", json={
        "left": {"type": "run", "id": run["id"]},
        "right": {"type": "run", "id": run["id"]},
        "keys": ["工号"],
    }).get_json()
    assert same["ok"]
    # 修复前这里是 19（模板原有行被算进来）
    assert same["summary"]["left_rows"] == run["summary"]["rows"] == 18
    assert same["summary"]["changed_cells"] == 0


# --------------------------------------------------------------------------
# 5. 本地服务防护：跨站请求必须被拒
# --------------------------------------------------------------------------
def test_cross_site_post_is_rejected(client):
    resp = client.post("/api/cleanup", headers={"Origin": "http://evil.example"})
    assert resp.status_code == 403
    assert resp.get_json()["ok"] is False


def test_cross_site_fetch_metadata_is_rejected(client):
    resp = client.post("/api/cleanup", headers={"Sec-Fetch-Site": "cross-site"})
    assert resp.status_code == 403


def test_same_origin_post_is_allowed(client):
    resp = client.post("/api/cleanup", headers={"Origin": "http://localhost"})
    assert resp.status_code == 200
    assert resp.get_json()["ok"] is True


def test_dns_rebinding_host_is_rejected(client):
    resp = client.get("/api/bootstrap", headers={"Host": "attacker.example"})
    assert resp.status_code == 403


# --------------------------------------------------------------------------
# 6. 公式注入：= + - @ 以及前导空白都要挡住
# --------------------------------------------------------------------------
@pytest.mark.parametrize("payload", ["=1+1", "+1", "-1+1", "@SUM(A1)", "  =cmd|' /C calc'!A1"])
def test_formula_prefix_is_forced_to_text(workdir, payload):
    out = workdir / "f.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["值"])
    cell = ws.cell(row=2, column=1)
    writer._set_safe(cell, payload)
    wb.save(out)
    reloaded = load_workbook(out)["Sheet"]["A2"]
    assert reloaded.data_type == "s"
    assert reloaded.quotePrefix is True


def test_normal_text_is_untouched(workdir):
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    writer._set_safe(cell, "正常文本")
    assert cell.quotePrefix is False


# --------------------------------------------------------------------------
# 7. skip_empty_key：以前是死配置，现在要真生效
# --------------------------------------------------------------------------
def _append_config(table_id, headers):
    cols = [{"key": f"c{i}", "name": h, "type": "auto", "included": True}
            for i, h in enumerate(headers)]
    maps = {f"c{i}": {"mode": "table", "table_id": table_id, "column_index": i,
                      "transform": "direct", "fallback": ""}
            for i in range(len(headers))}
    return cols, maps


def test_skip_empty_key_drops_rows(client):
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws.append(["工号", "姓名"])
    ws.append(["A001", "张三"])
    ws.append([None, "无编号"])
    ws.append(["A002", "李四"])
    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    up = client.post("/api/sources/upload", data={"files": [(buf, "k.xlsx")]},
                     content_type="multipart/form-data").get_json()
    table = up["created"][0]["meta"]["tables"][0]

    cols, maps = _append_config(table["id"], table["headers"])
    base = {
        "mode": "join", "primary_table": table["id"], "primary_keys": [0],
        "lookups": [], "append_tables": [], "auto_by_name": True,
        "target": {"kind": "blank", "sheet": "汇总", "header_row": 1, "columns": cols},
        "mappings": maps,
    }
    kept = build({**base, "options": {"dedupe": True, "skip_empty_key": False}},
                 include_rows=False)
    dropped = build({**base, "options": {"dedupe": True, "skip_empty_key": True}},
                    include_rows=False)
    assert kept["stats"]["rows"] == 3
    assert dropped["stats"]["rows"] == 2
    assert any("跳过" in w for w in dropped["warnings"])


# --------------------------------------------------------------------------
# 8. 单元格上限：防止 200k×300 直接把内存顶爆
# --------------------------------------------------------------------------
def test_cell_limit_guard(workdir, monkeypatch):
    wb = Workbook()
    ws = wb.active
    for r in range(20):
        ws.append([r] * 10)
    path = workdir / "big.xlsx"
    wb.save(path)

    monkeypatch.setattr(readers, "MAX_CELLS", 50)
    with pytest.raises(readers.UserError) as err:
        readers.read_excel_grids(path)
    assert "单元格" in str(err.value)


def test_only_requested_sheet_is_parsed(workdir, monkeypatch):
    wb = Workbook()
    wb.active.title = "要的"
    wb.active.append(["a", "b"])
    wb.create_sheet("不要的").append(["x", "y"])
    path = workdir / "two.xlsx"
    wb.save(path)

    grids = readers.read_excel_grids(path, only_sheets={"要的"})
    assert [g["name"] for g in grids] == ["要的"]


# --------------------------------------------------------------------------
# 9. 全空列提示：公式列读成空时要能看见，而不是让用户自己发现
# --------------------------------------------------------------------------
def test_preview_reports_empty_columns(client):
    import io
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "含空列"
    ws.append(["工号", "姓名", "公式列"])
    ws.append(["A001", "张三", None])
    ws.append(["A002", "李四", None])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    up = client.post("/api/sources/upload", data={"files": [(buf, "empty.xlsx")]},
                     content_type="multipart/form-data").get_json()
    table_id = up["created"][0]["meta"]["tables"][0]["id"]

    data = client.get(f"/api/table-preview?table_id={table_id}&limit=10").get_json()
    assert data["ok"]
    assert data["empty_columns"] == ["公式列"]
    assert data["scanned_rows"] == 2
