"""生成自测用的办公样例文件（员工主表、考勤、工资、模板、邮件）。"""
from __future__ import annotations

import csv
import datetime as dt
import email.message
import email.policy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent

employees = [
    ("E001", "陈思远", "研发部", "后端工程师", "2021-03-15", "13800001001"),
    ("E002", "林嘉禾", "市场部", "市场经理", "2019-07-01", "13800001002"),
    ("E003", "周雨桐", "财务部", "会计", "2022-11-08", "13800001003"),
    ("E004", "许明哲", "研发部", "前端工程师", "2020-05-20", "13800001004"),
    ("E005", "苏晚晴", "人力部", "招聘主管", "2018-09-12", "13800001005"),
    ("E006", "顾言", "研发部", "测试工程师", "2023-02-06", "13800001006"),
    ("E007", "沈知夏", "市场部", "品牌专员", "2023-06-19", "13800001007"),
    ("E008", "江屿", "研发部", "算法工程师", "2021-10-11", "13800001008"),
    ("E009", "温以宁", "财务部", "出纳", "2022-04-25", "13800001009"),
    ("E010", "陆星河", "市场部", "渠道经理", "2017-12-04", "13800001010"),
    ("E011", "唐棠", "人力部", "薪酬专员", "2020-08-17", "13800001011"),
    ("E012", "程砚", "研发部", "产品经理", "2019-03-28", "13800001012"),
    ("E013", "宋清和", "财务部", "财务主管", "2016-06-13", "13800001013"),
    ("E014", "叶知秋", "市场部", "内容运营", "2024-01-09", "13800001014"),
    ("E015", "韩叙", "研发部", "运维工程师", "2021-09-01", "13800001015"),
    ("E016", "白露", "人力部", "培训专员", "2023-11-21", "13800001016"),
    ("E017", "季沉", "研发部", "数据工程师", "2022-07-05", "13800001017"),
    ("E018", "柳青", "财务部", "税务专员", "2020-02-24", "13800001018"),
]

attendance = [
    ("E001", 21, 0, 6), ("E002", 20, 1, 3), ("E003", 22, 0, 0), ("E004", 19, 2, 8),
    ("E005", 21, 1, 2), ("E006", 22, 0, 4), ("E007", 18, 3, 0), ("E008", 21, 0, 10),
    ("E009", 22, 0, 0), ("E010", 20, 2, 5), ("E011", 21, 0, 1), ("E012", 22, 0, 7),
    ("E013", 19, 1, 2), ("E014", 21, 0, 3), ("E015", 22, 0, 9), ("E016", 20, 1, 0),
    ("E017", 21, 0, 6), ("E018", 22, 0, 2), ("E099", 22, 0, 0),
]

salary = [
    ("E001", 23800, 3800, 2856), ("E002", 26500, 4200, 3180), ("E003", 15600, 2500, 1872),
    ("E004", 21500, 3440, 2580), ("E005", 19800, 3170, 2376), ("E006", 17200, 2750, 2064),
    ("E007", 14800, 2370, 1776), ("E008", 28600, 4580, 3432), ("E009", 13200, 2110, 1584),
    ("E010", 24200, 3870, 2904), ("E011", 16800, 2690, 2016), ("E012", 27800, 4450, 3336),
    ("E013", 22600, 3620, 2712), ("E014", 12600, 2020, 1512), ("E015", 20500, 3280, 2460),
    ("E016", 13900, 2220, 1668), ("E017", 25200, 4030, 3024),
]

thin = Side(style="thin", color="DDE6E1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def header_row(ws, row, values, fill="0F6B5C"):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border


def make_employees():
    wb = Workbook()
    ws = wb.active
    ws.title = "在职人员"
    ws["A1"] = "2026 年 10 月员工主数据"
    ws["A1"].font = Font(name="微软雅黑", size=15, bold=True, color="0F6B5C")
    ws.merge_cells("A1:G1")
    header_row(ws, 2, ["工号", "姓名", "部门", "岗位", "入职日期", "手机号", "员工状态"])
    for r, emp in enumerate(employees, start=3):
        for c, v in enumerate(emp, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = border
        ws.cell(row=r, column=7, value="在职")
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
    widths = [10, 12, 12, 16, 14, 16, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    wb.save(OUT / "员工主表.xlsx")


def make_attendance():
    wb = Workbook()
    ws = wb.active
    ws.title = "10月考勤"
    ws["A1"] = "考勤统计（2026-10）"
    ws["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    header_row(ws, 3, ["员工编号", "出勤天数", "迟到次数", "加班小时"])
    for r, row in enumerate(attendance, start=4):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = border
    for i, w in enumerate([12, 12, 12, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    wb.save(OUT / "考勤统计.xlsx")


def make_salary():
    with open(OUT / "工资表.csv", "w", encoding="gbk", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["工号", "应发工资", "社保", "公积金"])
        writer.writerows(salary)


def make_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "人员汇总"
    ws["A1"] = "2026 年 10 月人员信息汇总表"
    ws["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="0F6B5C")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A1:H1")
    ws.row_dimensions[1].height = 36
    header_row(ws, 2, ["工号", "姓名", "部门", "岗位", "出勤天数", "迟到次数", "应发工资", "备注"])
    ws.row_dimensions[2].height = 28
    # 一行旧数据，用来验证“覆盖数据区”
    for c, v in enumerate(["E000", "旧数据", "待删除", "", 0, 0, 0, ""], start=1):
        cell = ws.cell(row=3, column=c, value=v)
        cell.font = Font(name="微软雅黑", size=10, color="8A9690")
        cell.border = border
        cell.fill = PatternFill("solid", fgColor="F4F7F5")
    for i, w in enumerate([10, 12, 14, 16, 12, 12, 14, 24], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    wb.save(OUT / "人员汇总模板.xlsx")


def make_email():
    msg = email.message.EmailMessage(policy=email.policy.default)
    msg["Subject"] = "10月人员通讯录更新（含附件）"
    msg["From"] = "人事共享中心 <hr@example.com>"
    msg["To"] = "office@example.com"
    msg["Date"] = dt.datetime(2026, 10, 15, 9, 30).strftime("%a, %d %b %Y %H:%M:%S +0800")
    html = """<html><body>
    <p>各位同事：</p>
    <p>附件为更新后的办公地点信息，正文为邮箱补充表，请合并使用。</p>
    <table border="1" cellspacing="0" cellpadding="4">
      <tr><th>工号</th><th>企业邮箱</th><th>办公地点</th></tr>
      <tr><td>E001</td><td>chensiyuan@example.com</td><td>A座 5F</td></tr>
      <tr><td>E002</td><td>linjiahe@example.com</td><td>B座 3F</td></tr>
      <tr><td>E003</td><td>zhouyutong@example.com</td><td>A座 6F</td></tr>
      <tr><td>E004</td><td>xumingzhe@example.com</td><td>A座 5F</td></tr>
      <tr><td>E005</td><td>suwanqing@example.com</td><td>B座 2F</td></tr>
    </table>
    <p>人力资源部</p>
    </body></html>"""
    msg.set_content("请查看 HTML 版本邮件。")
    msg.add_alternative(html, subtype="html")
    wb = Workbook()
    ws = wb.active
    ws.title = "办公地点补充"
    header_row(ws, 1, ["工号", "办公地点", "分机号"])
    extra = [("E006", "A座 5F", "8606"), ("E007", "B座 3F", "8307"), ("E008", "A座 5F", "8608")]
    for r, row in enumerate(extra, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    for i, w in enumerate([10, 14, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    import io
    buf = io.BytesIO()
    wb.save(buf)
    msg.add_attachment(buf.getvalue(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       filename="办公地点补充.xlsx")
    (OUT / "通讯录更新.eml").write_bytes(msg.as_bytes())


if __name__ == "__main__":
    make_employees()
    make_attendance()
    make_salary()
    make_template()
    make_email()
    print("样例文件已生成：", OUT)
