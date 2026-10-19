"""生成《表格汇总助手 · 使用说明》DOCX。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets" / "derived"
ICON = ROOT.parent / "packaging" / "huiyi-icon.png"

BRAND = RGBColor(0x0E, 0x6B, 0x5C)
BRAND_DARK = RGBColor(0x0A, 0x57, 0x4A)
INK = RGBColor(0x16, 0x21, 0x1C)
INK_SOFT = RGBColor(0x3D, 0x4B, 0x44)
MUTED = RGBColor(0x72, 0x80, 0x78)
FAINT = RGBColor(0x98, 0xA4, 0x9D)
GOLD = RGBColor(0x9A, 0x71, 0x29)
RED = RGBColor(0xB4, 0x47, 0x3D)
LIGHT = "EAF5F0"
LIGHT_GOLD = "FBF2E1"
LIGHT_GREY = "F5F8F6"
LINE = "DDE6E1"

FONT = "Microsoft YaHei"
EA_FONT = "微软雅黑"
BODY = 10.5
USABLE = 16.7


def set_run(run, size=BODY, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    return run


def set_para(p, before=0, after=5, line=1.4, align=None, keep=False, left=0, first=0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    if left:
        pf.left_indent = Cm(left)
    if first:
        pf.first_line_indent = Cm(first)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_margins(cell, top=80, bottom=80, left=110, right=110):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def cell_border(cell, edge, color=LINE, sz=4, val="single"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    el = borders.find(qn(f"w:{edge}"))
    if el is None:
        el = OxmlElement(f"w:{edge}")
        borders.append(el)
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for i, width in enumerate(widths):
        table.columns[i].width = Cm(width)
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Cm(width)


def table(headers, rows, widths, header_fill="0E6B5C", zebra=True, font_size=9.6,
          first_bold=False, align_center=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_width(t, widths)
    t.rows[0].height = Cm(0.9)
    trpr0 = t.rows[0]._tr.get_or_add_trPr()
    trpr0.append(OxmlElement("w:tblHeader"))
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, header_fill)
        cell_margins(cell)
        for edge in ("top", "left", "bottom", "right"):
            cell_border(cell, edge, color=header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=0, line=1.15)
        set_run(p.add_run(text), size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for r_idx, row in enumerate(rows):
        t.rows[r_idx + 1].height = Cm(0.8)
        trpr = t.rows[r_idx + 1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        for c_idx, text in enumerate(row):
            cell = t.rows[r_idx + 1].cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            fill = LIGHT_GREY if zebra and r_idx % 2 == 0 else "FFFFFF"
            shade(cell, fill)
            for edge in ("top", "left", "bottom", "right"):
                cell_border(cell, edge, color=LINE, sz=4)
            p = cell.paragraphs[0]
            set_para(p, after=0, line=1.25,
                     align=WD_ALIGN_PARAGRAPH.CENTER if (align_center and c_idx in align_center) else None)
            set_run(p.add_run(str(text)), size=font_size, bold=(first_bold and c_idx == 0),
                    color=INK if not (first_bold and c_idx == 0) else INK)
    return t


def callout(title, text, kind="tip", icon="●"):
    fill = LIGHT if kind == "tip" else LIGHT_GOLD if kind == "warn" else LIGHT_GREY
    accent = BRAND if kind == "tip" else GOLD if kind == "warn" else RGBColor(0x55, 0x63, 0x5D)
    accent_hex = "0E6B5C" if kind == "tip" else "9A7129" if kind == "warn" else "55635D"
    t = doc.add_table(rows=1, cols=1)
    set_table_width(t, [USABLE])
    trpr = t.rows[0]._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trpr.append(cant)
    cell = t.rows[0].cells[0]
    shade(cell, fill)
    cell_margins(cell, top=120, bottom=120, left=170, right=170)
    for edge in ("top", "bottom", "right"):
        cell_border(cell, edge, color=fill)
    cell_border(cell, "left", color=accent_hex, sz=18)
    p = cell.paragraphs[0]
    set_para(p, after=0, line=1.35)
    set_run(p.add_run(f"{title}："), size=9.8, bold=True, color=accent)
    set_run(p.add_run(text), size=9.8, color=INK_SOFT)
    spacer(0.18)


def spacer(cm=0.18):
    p = doc.add_paragraph()
    set_para(p, before=0, after=0, line=1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("")
    set_run(run, size=1)


def H1(text, number=""):
    p = doc.add_paragraph()
    set_para(p, before=6, after=6, line=1.2, keep=True)
    if number:
        set_run(p.add_run(number + "  "), size=20, bold=True, color=RGBColor(0xC9, 0xE3, 0xD9))
    set_run(p.add_run(text), size=19, bold=True, color=BRAND_DARK)
    # bottom rule
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "9AC7B8")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(12)


def H2(text):
    p = doc.add_paragraph()
    set_para(p, before=11, after=4, line=1.25, keep=True)
    set_run(p.add_run(text), size=12.5, bold=True, color=INK)
    return p


def P(text, size=BODY, color=INK_SOFT, before=0, after=5, bold=False):
    p = doc.add_paragraph()
    set_para(p, before=before, after=after, line=1.45)
    set_run(p.add_run(text), size=size, color=color, bold=bold)
    return p


def rich(parts, before=0, after=5, size=BODY, indent=0):
    p = doc.add_paragraph()
    set_para(p, before=before, after=after, line=1.45, left=indent)
    for text, bold, color in parts:
        set_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def bullets(items, size=BODY):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_para(p, before=0, after=3, line=1.4, left=0.55, first=-0.35)
        if isinstance(item, tuple):
            lead, rest = item
            set_run(p.add_run(lead), size=size, bold=True, color=INK)
            set_run(p.add_run(rest), size=size, color=INK_SOFT)
        else:
            set_run(p.add_run(item), size=size, color=INK_SOFT)


def step(num, title, desc):
    p = doc.add_paragraph()
    set_para(p, before=5, after=2, line=1.35, left=0.75, first=-0.75)
    set_run(p.add_run(f"{num}"), size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    # 简单给数字加深色底不方便，改用彩色圆圈字符
    p._p.remove(p.runs[0]._r)
    set_run(p.add_run("● "), size=8, color=BRAND)
    set_run(p.add_run(title), size=10.8, bold=True, color=INK)
    if desc:
        p2 = doc.add_paragraph()
        set_para(p2, before=0, after=4, line=1.4, left=0.75)
        set_run(p2.add_run(desc), size=9.8, color=INK_SOFT)


def figure(path, caption, width=16.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, before=4, after=2, line=1, keep=True)
    p.add_run().add_picture(str(path), width=Cm(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(cap, before=0, after=8, line=1.2)
    set_run(cap.add_run(caption), size=8.8, color=FAINT)


def add_page_number(paragraph):
    set_run(paragraph.add_run("第 "), size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    rpr.append(rfonts)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "17"); rpr.append(sz)
    color = OxmlElement("w:color"); color.set(qn("w:val"), "728078"); rpr.append(color)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = "1"; run.append(t)
    fld.append(run)
    paragraph._p.append(fld)
    set_run(paragraph.add_run(" 页"), size=8.5, color=MUTED)


def setup_document():
    global doc
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0.9)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    normal.font.color.rgb = INK_SOFT
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), EA_FONT)

    # header
    header = sec.header
    hp = header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(USABLE), WD_TAB_ALIGNMENT.RIGHT)
    set_run(hp.add_run("表格汇总助手 · 使用说明"), size=8.5, color=MUTED)
    set_run(hp.add_run("\t内部使用 · v0.1.0"), size=8.5, color=FAINT)
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "3"); bottom.set(qn("w:color"), "DDE6E1")
    pBdr.append(bottom); pPr.append(pBdr)

    # footer
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(USABLE), WD_TAB_ALIGNMENT.RIGHT)
    set_run(fp.add_run("把多份表格，汇成一张表"), size=8.5, color=FAINT)
    set_run(fp.add_run("\t"), size=8.5)
    add_page_number(fp)


def cover():
    # 顶部品牌
    p = doc.add_paragraph()
    set_para(p, before=2, after=2, line=1)
    p.add_run().add_picture(str(ICON), width=Cm(2.1))
    p2 = doc.add_paragraph()
    set_para(p2, before=0, after=0, line=1.1)
    set_run(p2.add_run("表 格 汇 总 助 手"), size=8.5, bold=True, color=BRAND)
    p3 = doc.add_paragraph()
    set_para(p3, before=20, after=2, line=1.1)
    set_run(p3.add_run("表格汇总助手"), size=30, bold=True, color=BRAND_DARK)
    p4 = doc.add_paragraph()
    set_para(p4, before=2, after=0, line=1.3)
    set_run(p4.add_run("办公室表格汇总 · 使用说明"), size=16, bold=False, color=INK)
    p5 = doc.add_paragraph()
    set_para(p5, before=5, after=3, line=1.3)
    set_run(p5.add_run("把多份表格，汇成一张表"), size=12, bold=True, color=GOLD)
    p5.paragraph_format.space_after = Pt(10)
    # rule
    rule = doc.add_paragraph()
    set_para(rule, before=0, after=10, line=1)
    pPr = rule._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "0E6B5C")
    pBdr.append(bottom); pPr.append(pBdr)

    P("把 Excel、CSV、Outlook 邮件和附件里的数据，按你的规则汇总到一张指定表格里。每一列都能单独指定来源，生成结果保留模板格式，还能和上一次结果逐行对比。",
      size=10.5, color=INK_SOFT, after=10)
    figure(ASSETS / "cover.png", "工作台首页：上传文件、建立方案、生成与对比", width=16.5)
    spacer(0.15)

    info = doc.add_table(rows=3, cols=3)
    set_table_width(info, [5.2, 5.7, 5.8])
    info_rows = [
        ("适用对象", "办公室文员 / 行政 / 人事 / 财务", "运行方式", "安装在本机，浏览器打开"),
        ("数据方式", "不联网，不上传云端", "版本", "v0.1.0"),
        ("建议先看", "第 2 页「先看这一页」", "练习文件", "samples 文件夹"),
    ]
    # 3 行 3 列不够，改为纵向信息块
    for row in info.rows:
        for cell in row.cells:
            cell.text = ""
    # 重新构建成两列表，避免复杂
    info._tbl.getparent().remove(info._tbl)
    info = doc.add_table(rows=3, cols=2)
    set_table_width(info, [4.0, 12.4])
    for i, (k, v) in enumerate([("适用对象", "办公室文员 / 行政 / 人事 / 财务"),
                                 ("使用方式", "安装在本机，浏览器打开；不联网、不上传云端"),
                                 ("版本", "v0.1.0 · 2026 年 10 月")]):
        c0, c1 = info.rows[i].cells
        shade(c0, LIGHT); shade(c1, "FFFFFF")
        cell_margins(c0); cell_margins(c1)
        for edge in ("top", "bottom", "left", "right"):
            cell_border(c0, edge, color=LINE); cell_border(c1, edge, color=LINE)
        p0 = c0.paragraphs[0]; set_para(p0, after=0, line=1.2)
        set_run(p0.add_run(k), size=9, bold=True, color=BRAND_DARK)
        p1 = c1.paragraphs[0]; set_para(p1, after=0, line=1.2)
        set_run(p1.add_run(v), size=9, color=INK_SOFT)
    doc.add_page_break()


def page_two():
    H1("先看这一页", "开始")
    P("不用担心步骤多。你只需要记住：先把文件放进来，再告诉工具“每一列从哪里取数”，最后点生成。中间不需要写公式，也不需要会 Excel 宏。", after=8)

    H2("整个流程：5 个动作")
    table(
        ["顺序", "你要做的事", "在哪个页面"],
        [
            ("1", "上传 Excel、CSV、邮件或附件", "数据源"),
            ("2", "新建一个汇总方案", "汇总方案"),
            ("3", "选择主表、主键，添加要补充的表", "方案编辑 · 数据来源"),
            ("4", "上传目标模板，给每一列选来源", "方案编辑 · 目标表 / 字段映射"),
            ("5", "预览、生成、下载；需要时再做对比", "方案编辑 / 生成记录 / 数据对比"),
        ],
        [1.2, 10.6, 4.7],
        align_center=[0],
    )
    spacer(0.25)

    H2("先认识 8 个词")
    table(
        ["你会看到的词", "用大白话解释"],
        [
            ("数据源", "你上传的原始文件。可以是 Excel、CSV、邮件或邮件附件。"),
            ("主表", "决定最后表格有多少行的那张表，通常是人名/订单最全的那张。"),
            ("主键", "两张表都有的、能一一对上的那一列，比如工号、订单号。"),
            ("关联表", "给主表补充字段的表。它只补数据，不会增加行数。"),
            ("目标表 / 模板", "最后要填写的那张 Excel。可以上传公司现有模板。"),
            ("列来源", "最后表里的每一列，从哪张表、哪个字段取数。"),
            ("缺省值", "找不到数据时，默认填入什么。"),
            ("覆盖 / 追加", "覆盖 = 清掉模板旧数据再填；追加 = 接在已有数据后面。"),
        ],
        [3.6, 12.6],
        first_bold=True,
    )
    spacer(0.22)
    callout("开始前，最好准备三样东西",
            "① 原始表格；② 最后要填的 Excel 模板（没有也可以新建）；③ 一个能对上号的列（工号、订单号、身份证号等）。",
            kind="tip")
    callout("第一次使用，先用示例文件练手",
            "samples 文件夹里有员工表、考勤表、工资 CSV、汇总模板和邮件。照本说明走一遍，熟悉后再换成自己的文件。",
            kind="example")

def chapter_upload():
    H1("把文件放进来", "第一步")
    P("打开程序后，浏览器会自动进入工作台。左侧「数据源」就是放文件的地方。上传只是把文件接入工具，不会改动原文件。", after=6)

    H2("1. 上传文件")
    bullets([
        ("拖进来：", "把文件直接拖到虚线框里。"),
        ("点按钮：", "点「上传数据」或「浏览文件」，一次可以选多个。"),
        ("等几秒：", "文件越大，解析时间越长；进度条消失就完成了。"),
    ])
    figure(ASSETS / "sources.png", "图 1 · 数据源页面：上方是上传区，下方是已经接入的文件", width=16.3)

    H2("2. 看一眼文件卡片")
    P("每张卡片会显示文件名、类型、大小、工作表/数据表和行数。点表格标签或右侧的「眼睛」图标，可以预览前 30 行。", after=4)
    callout("邮件文件不用手动拆附件",
            "上传 .msg 或 .eml 后，正文里的表格会自动变成数据表，附件会展开在邮件卡片下方，可以单独预览。",
            kind="tip")
    P("支持的文件类型：Excel（.xlsx / .xlsm / .xls）、CSV / TSV、Outlook 邮件（.msg / .oft）、.eml。暂不支持 .pst / .ost，请先在 Outlook 里另存为 .msg。",
      size=9.6, color=MUTED, before=4)

    H2("3. 检查表头（很重要）")
    P("工具会自动找“像表头”的那一行，但有些表格前面有标题、说明或空行，自动识别可能不准。表头不对，后面所有列都会错位，所以上传后一定先点开预览。", after=4)
    figure(ASSETS / "source-preview.png", "图 2 · 预览弹窗：右上角可以修改“表头所在行”", width=15.0)
    bullets([
        ("表头行：", "一行列名的位置，比如第 2 行。大标题占了一行时，表头往往在第 2 或第 3 行。"),
        ("列名：", "确认列名和下面的数据对得上。"),
        ("行数：", "和原文件大致对一下，少了可能是表头识别错、或者中间有空行。"),
    ])
    callout("修改表头行后，记得点「应用」",
            "修改后，系统会重新按新的表头读取列。如果方案里已经做过映射，请打开方案再看一眼，必要时重新选一下列。",
            kind="warn")

def chapter_profile():
    H1("建一个汇总方案", "第二步")
    P("“方案”可以理解成一张固定的工作说明书：用哪张主表、按什么主键、补哪些表、每一列从哪来。配置一次，下个月只要换文件，点一下就能重新生成。", after=6)

    H2("1. 新建方案")
    figure(ASSETS / "create-profile.png", "图 3 · 新建汇总方案：填名称、选汇总方式、可选模板", width=14.5)
    bullets([
        ("方案名称：", "写清楚用途和周期，例如“10 月人员信息汇总”。"),
        ("汇总方式：", "大多数人员、订单类汇总选「主表关联补齐」；多个月份/部门的同构表选「多表纵向合并」。"),
        ("模板：", "如果最后要填公司固定格式的 Excel，选择已有模板；也可以先不选，进方案后再上传。"),
    ])
    callout("先不用急着把所有东西都填完",
            "方案可以边配边保存。页面会自动保存，右侧「生成前检查」会提示还缺什么。",
            kind="tip")

    H2("2. 选主表、主键和关联表")
    P("以示例为例：员工主表有 18 个人，是“主表”；考勤、工资、通讯录是“关联表”。主表决定最后有 18 行，关联表只把字段补进来。", after=4)
    figure(ASSETS / "profile-source.png", "图 4 · 数据来源：主表、主键和关联表都在这一个区域设置", width=16.0)
    table(
        ["要选/要填", "怎么理解", "示例"],
        [
            ("主表", "最后表格的行数、顺序以它为准", "员工主表"),
            ("主键", "用来和别的表“对上号”的列，可选多个组成组合主键", "工号"),
            ("关联表", "补充字段的表；按主键匹配，不增加行", "考勤表、工资表、通讯录"),
            ("关联字段", "关联表里用来对上主键的列；名字可以不同", "员工编号 ↔ 工号"),
        ],
        [3.0, 9.2, 3.8],
        first_bold=True,
    )
    spacer(0.2)
    callout("举个例子",
            "员工主表有“工号”，考勤表里叫“员工编号”。在关联表设置里，把主表字段「工号」对应到考勤表字段「员工编号」，就能对上。",
            kind="example")
    callout("对不上时不要硬生成",
            "如果大量关联表显示空白，先检查：两边是不是一个有空格、一个是纯数字；工号有没有重复；主键是不是选错了列。",
            kind="warn")

    H2("3. 选择目标表")
    P("在「目标表」区域上传最后要填的 Excel。工具会读取表头，并保留模板的列宽、颜色、合并单元格和未参与写入的公式列。", after=4)
    table(
        ["选项", "什么时候用"],
        [
            ("覆盖模板数据区", "模板是空白表，或每个月都重新填一遍。旧数据会被清掉。"),
            ("追加到已有数据之后", "模板里已经有历史数据，想把新结果接在后面。"),
            ("纳入 / 不纳入", "不纳入的列保持模板原样，适合公式列、签字列、手工备注列。"),
        ],
        [5.2, 11.0],
        first_bold=True,
    )
    spacer(0.18)
    figure(ASSETS / "template-modal.png", "图 5 · 目标模板：选择已有模板，或上传一个新的 Excel", width=14.8)
    callout("生成不会修改你的原始模板",
            "每次生成都会复制一份模板再写入，原模板保持不动。第一次使用复杂模板时，建议先生成一份，打开核对格式。",
            kind="tip")

def chapter_mapping():
    doc.add_page_break()
    H1("每一列从哪里来", "第三步")
    P("这是整个工具最核心的一步。目标表里的每一列，都要告诉工具：从哪张表、哪个字段取数；找不到时填什么。", after=6)
    figure(ASSETS / "profile-mapping.png", "图 6 · 字段映射：一行对应一个目标列", width=16.2)

    H2("每行从左到右看")
    table(
        ["位置", "你看到的", "要做的事"],
        [
            ("第 1 列", "目标列", "最后表格里的列名，来自模板或手动添加。"),
            ("第 2 列", "数据来源", "选主表、某张关联表，或「固定值」「留空」。"),
            ("第 3 列", "来源字段", "选那张表里的具体列名。"),
            ("第 4 列", "处理", "一般保持原值；需要时去空格、转数值、统一日期。"),
            ("第 5 列", "缺省值", "找不到或为空时填什么；可以留空。"),
            ("最右", "绿色勾 / 黄色叹号", "绿色表示已配置；黄色表示缺少来源或来源失效。"),
        ],
        [2.2, 4.0, 10.3],
        first_bold=True,
    )
    spacer(0.2)

    H2("推荐的配置顺序")
    bullets([
        ("先自动匹配：", "点「按列名匹配」，工具会把目标列和来源表的同名列自动连起来。"),
        ("再逐行检查：", "重点看没有变绿的列，手动选来源字段。"),
        ("固定值：", "整列都写同一个内容时用，例如“备注 = 2026-11”。"),
        ("留空：", "这一列暂时不输出数据；模板里未纳入的列会保持原样。"),
    ])
    callout("“去首尾空格”很有用",
            "Excel 里看起来一样的“E001”和“E001 ”，实际可能对不上。主键或工号列建议选「去除首尾空格」。",
            kind="tip")
    callout("类型怎么选",
            "不确定就选「自动」。需要写日期格式时选「日期」；需要计算金额、数量时选「数值」；编号建议选「文本」，避免前导 0 丢失。",
            kind="example")

    H2("做完映射，用 4 个问题自查")
    table(
        ["检查问题", "合格标准"],
        [
            ("每个目标列都有来源吗？", "全部是绿色勾，没有黄色叹号。"),
            ("工号/订单号处理了吗？", "主键列建议选「去除首尾空格」，避免空格导致对不上。"),
            ("找不到数据时填什么？", "缺省值或固定值已设置；不需要填的可以留空。"),
            ("不写的列会保留吗？", "模板里未纳入的列保持原样，适合公式列、签字列。"),
        ],
        [6.0, 10.2],
        first_bold=True,
    )


def chapter_generate():
    doc.add_page_break()
    H1("预览、生成、下载", "第四步")
    P("配置完成后，不要直接生成。先预览前 50 行，确认行数、列名和空值情况，再生成完整文件。", after=6)
    figure(ASSETS / "preview.png", "图 7 · 预览：上方是统计，下方是实际汇总结果", width=15.4)

    H2("预览时看三件事")
    table(
        ["看什么", "正常情况", "不正常时怎么办"],
        [
            ("行数", "和主表行数一致（或接近）", "检查主键、去重开关、主表是否选错"),
            ("空值", "只有少数格子为空", "检查关联表、主键、缺省值"),
            ("列名", "和目标模板一致", "检查目标表列、模板表头行"),
        ],
        [3.2, 6.2, 7.1],
        first_bold=True,
    )
    spacer(0.2)
    callout("预览只显示前 50 行",
            "预览不会写入文件，也不会改动数据源。真正生成时会读取全部数据（单表上限 20 万行）。",
            kind="tip")

    doc.add_page_break()
    H2("生成与下载")
    bullets([
        ("点「生成汇总表」：", "工具会先自动保存方案，再生成 Excel。"),
        ("看结果弹窗：", "确认输出行数、填充率和输出工作表；有提示时展开看一眼。"),
        ("下载：", "点「下载结果」保存到电脑；也可以稍后到「生成记录」重新下载。"),
        ("文件位置：", "程序目录下的 data/runs 文件夹。建议下载后复制到自己的业务文件夹。"),
    ])
    figure(ASSETS / "runs.png", "图 8 · 生成记录：下载、查看详情、发起对比", width=16.0)
    callout("重要结果请及时下载",
            "工具会按设置只保留最近 30 条生成记录，超过后自动清理旧结果。重要文件请自己留一份。",
            kind="warn")

    H2("生成后建议")
    bullets([
        ("先打开看一眼：", "确认表头、列数、总行数和关键数据都对。"),
        ("检查格式：", "金额、日期、编号的显示格式是否符合业务习惯。"),
        ("另存一份：", "把结果复制到自己的业务文件夹，不要只留在生成目录。"),
        ("保留上期结果：", "下一次可以和它做对比，方便核对人员或订单变化。"),
    ])

def chapter_compare():
    doc.add_page_break()
    H1("对比两次结果", "第五步")
    P("月度汇总、交接核对时，可以用「数据对比」找出两次结果之间的新增、删除和修改。对比不会改动原文件。", after=6)
    figure(ASSETS / "compare.png", "图 9 · 数据对比：选左右两侧和主键，就能看到差异", width=16.2)

    H2("操作顺序")
    table(
        ["步骤", "做什么"],
        [
            ("1", "左侧选旧结果（基准），右侧选新结果；也可以选两张原始数据表。"),
            ("2", "选一个共同主键，例如「工号」；支持组合主键。"),
            ("3", "点「开始对比」。"),
            ("4", "看四个数字：新增、删除、修改、未变化。"),
            ("5", "在下方按类型筛选，或点「导出对比报告」保存 Excel。"),
        ],
        [1.2, 15.0],
        align_center=[0],
    )
    spacer(0.2)
    bullets([
        ("修改：", "同一主键的某个字段，左右值不同。"),
        ("新增 / 删除：", "只在一侧出现的主键；通常是人员入职/离职、订单增加/取消。"),
        ("字段差异：", "两侧列不完全一致时会单独提示，只比较共同列。"),
    ])
    callout("对比前先对齐主键",
            "如果左边用工号、右边用姓名，会出现大量“新增/删除”。请选择两侧都有的同一个列。",
            kind="warn")

def chapter_faq():
    H1("遇到问题先看这里", "常见问题")
    qa = [
        ("关联表大量空白、对不上",
         "先检查主键：两边是否都有值、有没有多余空格、是不是一边文本一边数字、有没有重复。可在主键列使用「去除首尾空格」；重复主键默认只保留第一行。"),
        ("生成后有空列 / 缺数据",
         "到「字段映射」看这一列是否变绿；检查来源字段是否选对、缺省值是否为空。再到预览看空值数。"),
        ("模板列没有出现，或列错位",
         "检查模板的“表头行”是否选对；在目标表区域点「调整」重新读取。复杂模板建议先用一份副本试生成。"),
        ("模板里的图表、图片、透视表不见了",
         "当前使用 openpyxl 写入，普通表头、列宽、合并单元格、公式列可以保留，但图表、图片、透视表等对象可能无法完整保留。重要模板请生成后人工核对。"),
        ("Outlook 的 .pst 文件能直接读吗？",
         "暂不支持。请在 Outlook 里把需要的邮件另存为 .msg，或导出为 .eml；正文表格和附件都会自动处理。"),
        ("文件太大 / 行数太多",
         "单个文件上限 100 MB，单表读取上限 20 万行。更大的表建议先拆分，或只保留需要汇总的列和行。"),
        ("页面打不开 / 提示端口占用",
         "确认启动程序的黑窗口还开着；在浏览器输入 http://127.0.0.1:8765。若端口被占用，启动窗口会显示新地址，以窗口为准。"),
        ("数据保存在哪里？怎么备份？",
         "程序目录下的 data 文件夹。重要方案、结果都在里面。备份时整个复制 data；升级程序时保留 data 即可。"),
        ("关闭黑窗口后页面不能用了",
         "这是正常的：黑窗口就是本地服务。重新双击启动程序即可；数据不会丢。"),
        ("多人能同时用吗？",
         "当前是单机版，装在自己电脑上使用；不提供多人同时在线和权限管理。不要把这个端口开放到不可信网络。"),
    ]
    for i, (q, a) in enumerate(qa, 1):
        p = doc.add_paragraph()
        set_para(p, before=6, after=1, line=1.35, left=0.7, first=-0.7)
        set_run(p.add_run(f"Q{i}. "), size=10.2, bold=True, color=BRAND)
        set_run(p.add_run(q), size=10.2, bold=True, color=INK)
        p2 = doc.add_paragraph()
        set_para(p2, before=0, after=5, line=1.4, left=0.7)
        set_run(p2.add_run(a), size=9.7, color=INK_SOFT)

def chapter_checklist():
    H1("最后一页：操作检查清单", "检查")
    P("每次正式生成前，照着打勾，能避免绝大多数问题。", after=8)
    items = [
        "文件都上传成功，文件名、工作表、行数正确",
        "每张表的表头行都对，预览里列名和数据能对上",
        "主表选的是行数基准表；主键在两边都有",
        "关联表的关联字段已选好，名字不同也能对应",
        "目标模板/工作表/表头行正确；覆盖或追加选对",
        "目标列都已配置，绿色的勾占多数",
        "主键列做了去空格；文本/数字格式一致",
        "预览的行数、空值、列名都检查过",
        "已生成并打开结果文件核对，重要文件已下载",
        "需要对比时，两侧主键一致，差异数量合理",
        "data 文件夹已备份（至少每月一次）",
    ]
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        set_para(p, before=2, after=2, line=1.45, left=0.9, first=-0.9)
        set_run(p.add_run("□  "), size=11, bold=True, color=BRAND)
        set_run(p.add_run(item), size=10.3, color=INK_SOFT)

    spacer(0.4)
    callout("遇到解决不了的问题，这样反馈最有效",
            "截一张图 + 说明在第几步 + 文件名/工作表名 + 错误提示原文。不要只发“用不了”，这样能最快定位。",
            kind="tip")
    callout("数据安全",
            "本工具在本机运行，不联网、不上传云端。不要把 data 文件夹发给无关人员；不要把服务端口开放到不可信网络。",
            kind="warn")



def build():
    setup_document()
    doc.core_properties.title = "表格汇总助手 · 使用说明"
    doc.core_properties.subject = "办公室多源表格汇总工具使用指南"
    doc.core_properties.author = "表格汇总助手"
    doc.core_properties.comments = "面向办公室文员的本地离线 Excel 汇总工具使用说明"
    cover()
    page_two()
    chapter_upload()
    chapter_profile()
    chapter_mapping()
    chapter_generate()
    chapter_compare()
    chapter_faq()
    chapter_checklist()
    out = ROOT / "表格汇总助手-使用说明.docx"
    doc.save(out)
    print("saved", out)


doc = None
if __name__ == "__main__":
    build()
